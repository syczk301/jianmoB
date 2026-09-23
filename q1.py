"""四问整理版：可直接阅读的合并源码，无隐藏代码包或动态执行源码。
公共物理方程位于q1.py；跨问共享函数通过普通Python模块导入。
数值模型、约束和原参数保持不变；默认命令仅整理已有结果。
"""
import os
import sys
import importlib
from pathlib import Path
os.environ.setdefault('OMP_NUM_THREADS', '1')
os.environ.setdefault('OPENBLAS_NUM_THREADS', '1')
if __name__ in ('__main__', '__mp_main__'):
    sys.modules.setdefault('q1', sys.modules[__name__])
def _question_module(name):
    return importlib.import_module(name)
def _project_root():
    return Path(os.environ.get('B_MODEL_WORKSPACE', str(Path(__file__).resolve().parent))).resolve()
def _project_storage():
    return Path(os.environ.get('B_MODEL_STORAGE', str(_project_root()/'work'))).resolve()

from dataclasses import dataclass, asdict, replace
import hashlib
import json
import numpy as np
from scipy.sparse import lil_matrix, csc_matrix
from dataclasses import dataclass
import time
from scipy.integrate import solve_ivp
import io
import platform
import sys
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
import pandas as pd
import openpyxl
from dataclasses import asdict
from matplotlib import font_manager, rcParams
from scipy.optimize import least_squares
from concurrent.futures import ProcessPoolExecutor
from scipy.stats import qmc
from scipy.linalg import solve_banded
from dataclasses import dataclass, asdict
from collections import deque
from scipy.linalg import eigh_tridiagonal
from dataclasses import replace
import matplotlib
import matplotlib.pyplot as plt
import subprocess
from concurrent.futures import ProcessPoolExecutor, as_completed
import os
import sys, json, time, argparse
import sys, json, argparse
import json, hashlib
from PIL import Image, ImageDraw, ImageFont
import csv, json, hashlib
import pytest

# ==================== config ====================
@dataclass(frozen=True)
class Config:
    cells: int = 1
    fixture: str = 'H1'
    basis: str = 'B'
    mesh: tuple = (3, 2, 3, 3, 3)
    board_n: int = 2
    end_n: int = 4
    shared_bp: bool = False
    water: str = 'W1'
    outlet: str = 'closed'
    thermal: str = 'phase'
    observation: str = 'MEA'
    j0: float = 0.01
    tau_b: float = 1.0
    tau_f: float = 1.0
    tau_m: float = 1.0
    mu_factor: float = 1.0
    ice_connectivity: float = 3.0
    h: float = 40.0
    contact_r: float = 0.0
    Rc: float = 1e-06
    lambda0: float = 3.0
    concentration: str = 'given'
    end_concentration_factor: float = 1.0
    cl_proton_loss: bool = False
    cathode_hydration_exponent: float = 0.0
    vapor_equilibrium: str = 'ice_reference'
    half_cl: bool = False
    T0: float = 253.15
    ambient: float = 253.15
    rtol: float = 1e-05
    max_step: float = 0.5
    voltage_limit: float = 0.3
    charge_limit: float = 20.0
    current_limit: float = 0.5
    success_margin: float = 0.01
    horizon: float = 300.0
    sample: float = 0.2
    control_prediction_horizon: float = 5.0
    seed: int = 20260923

    def changed(self, **kwargs):
        return replace(self, **kwargs)

    def to_dict(self):
        return asdict(self)

    @property
    def digest(self):
        return hashlib.sha256(json.dumps(asdict(self), sort_keys=True).encode()).hexdigest()[:16]

    def validate(self):
        assert self.cells in (1, 5)
        assert self.fixture in ('H1', 'H2')
        assert self.water in ('W0', 'W1')
        assert self.outlet in ('closed', 'drain')
        assert self.thermal in ('phase', 'effective')
        assert self.observation in ('MEA', 'unit')
        assert len(self.mesh) == 5 and min(self.mesh) >= 1
        assert self.j0 > 0 and min(self.tau_b, self.tau_f, self.tau_m) > 0
        assert self.cathode_hydration_exponent >= 0
        assert self.vapor_equilibrium in ('ice_reference', 'supercooled_liquid')

AREA = 0.0025

F = 96485.33212

R = 8.314462618

MW = 0.018

TM = 273.15

LV = 2500000.0

LF = 333600.0

P = 101325.0

YO2 = 0.233 / 0.032 / (0.233 / 0.032 + 0.767 / 0.028)

LAYER_NAMES = ('aGDL', 'aCL', 'PEM', 'cCL', 'cGDL')

THICKNESS = (0.00015, 3.4e-06, 1.2e-05, 1.13e-05, 0.00015)

POROSITY = (0.8, 0.3916, 0.0, 0.4207, 0.8)

DRY_C = (185 * 545, 970 * 240, 2150 * 1050, 970 * 240, 185 * 545)

K_SOLID = (0.3, 0.27, 0.24, 0.27, 0.3)

# ==================== properties ====================
def psat(T, liquid=False):
    c = np.clip(np.asarray(T) - TM, -90, 150)
    l = 611.21 * np.exp((18.678 - c / 234.5) * c / (257.14 + c))
    if liquid:
        return l
    i = 611.15 * np.exp((23.036 - c / 333.7) * c / (279.82 + c))
    return np.where(c >= 0, l, i)

def partition(w, ice, T, eps, liquid_reference=False):
    """Analytic saturation partition including the liquid-induced pore change.

    w=vapor+liquid, per total porous volume. Tiny negative trial states during
    Newton are extended only in constitutive evaluation, never projected into y.
    """
    wp = np.maximum(w, 0)
    rs = MW * psat(T, liquid=liquid_reference) / (R * T)
    sat = rs * (eps - ice / 920 - wp / 990) / (1 - rs / 990)
    mv = np.minimum(wp, np.maximum(0, sat))
    ml = wp - mv
    return (mv, ml, eps - ice / 920 - ml / 990)

def lambda_eq(aw):
    a = np.clip(aw, 0, 1)
    return 0.043 + 17.81 * a - 39.85 * a * a + 36 * a * a * a

def membrane_D(T, lam):
    return 1e-10 * np.exp(2416 * (1 / 303.15 - 1 / T)) * (2.563 - 0.33 * lam + 0.0264 * lam ** 2 - 0.000671 * lam ** 3)

def conductivity(T, lam):
    return (0.5139 * lam - 0.326) * np.exp(1268 * (1 / 303.15 - 1 / T))

def sigma(T):
    tau = 1 - np.asarray(T) / 647.096
    return 0.2358 * tau ** 1.256 * (1 - 0.625 * tau)

def liquid_mu(T, factor=1.0):
    tr = np.maximum(np.asarray(T), 253.15) / 300
    return factor * 1e-06 * (280.68 * tr ** (-1.9) + 511.45 * tr ** (-7.7) + 61.131 * tr ** (-19.6) + 0.45903 * tr ** (-40))

def hydraulic_pressure(ml, ice, T, eps, K, angle):
    """p_l-p_g; opposite sign to plan's p_c=p_g-p_l."""
    free = np.maximum(eps - ice / 920, 1e-12)
    s = np.clip(ml / 990 / free, 0, 1)
    J = 1.417 * s - 2.12 * s * s + 1.263 * s * s * s
    return -sigma(T) * np.cos(np.deg2rad(angle)) * np.sqrt(eps / K) * J

def water_enthalpies(T):
    dt = T - TM
    return (2000 * dt + LV, 4182 * dt, 2050 * dt - LF, 4182 * dt)

def harmonic_faces(dx, coeff):
    return 1 / (dx[:-1] / (2 * np.maximum(coeff[:-1], 1e-30)) + dx[1:] / (2 * np.maximum(coeff[1:], 1e-30)))

def liquid_flux(ml, ice, T, eps, K, angle, dx, factor=1.0, exponent=3.0, drain=False, side=0):
    p = hydraulic_pressure(ml, ice, T, eps, K, angle)
    saturation = np.clip(ml / (990 * np.maximum(eps - ice / 920, 1e-12)), 0, 1)
    connect = np.maximum(1 - ice / (920 * eps), 0) ** exponent
    mob = saturation ** 3 / liquid_mu(T, factor)
    f = np.zeros(len(dx) + 1)
    dp = p[:-1] - p[1:]
    f[1:-1] = 990 * harmonic_faces(dx, K * connect) * dp * np.where(dp >= 0, mob[:-1], mob[1:])
    if drain:
        b = 0 if side == 0 else -1
        sign = -1 if side == 0 else 1
        f[b] = sign * 990 * K[b] * connect[b] * mob[b] * max(p[b], 0) / (dx[b] / 2)
    return f

# ==================== geometry ====================
class Geometry:

    def __init__(self, cfg):
        self.cfg = cfg
        self.dx = np.concatenate([np.full(n, d / n) for d, n in zip(THICKNESS, cfg.mesh)])
        self.layer = np.repeat(np.arange(5), cfg.mesh)
        self.n = len(self.dx)
        self.x = np.cumsum(self.dx) - self.dx / 2
        self.eps = np.take(POROSITY, self.layer)
        self.omega = np.take([0, 0.3, 1, 0.3, 0], self.layer)
        self.bcap = self.omega * 2150 * 0.018
        self.cla = np.flatnonzero(self.layer == 1)
        self.clc = np.flatnonzero(self.layer == 3)
        self.pem = np.flatnonzero(self.layer == 2)
        self.bound = np.flatnonzero(self.omega > 0)
        self.anode = np.flatnonzero(self.layer < 2)
        self.cathode = np.flatnonzero(self.layer > 2)
        self.porous = self.eps > 0
        self.K = np.where(np.isin(self.layer, [1, 3]), 6.2e-13, 6.2e-12)
        self.angle = np.where(np.isin(self.layer, [1, 3]), 100.0, 110.0)
        self.im = np.zeros(self.n)
        self.im[self.cla] = (self.x[self.cla] - THICKNESS[0]) / THICKNESS[1]
        self.im[self.pem] = 1
        self.im[self.clc] = 1 - (self.x[self.clc] - sum(THICKNESS[:3])) / THICKNESS[3]
        self.gas_ref = np.where(self.layer < 2, 0.00011, 2.2e-05)
        self.vap_ref = np.where(self.layer < 2, 8.69e-05, 2.48e-05)
        self.gas_y = np.where(self.layer < 2, 1.0, 0.233 / 0.032 / (0.233 / 0.032 + 0.767 / 0.028))
        td, tc, tk, names, owners = ([], [], [], [], [])
        self.mea_indices = []
        self.heater_indices = []
        self.unit_indices = []

        def block(name, thickness, n, heatcap, k, owner):
            ix = np.arange(len(td), len(td) + n)
            td.extend([thickness / n] * n)
            tc.extend([heatcap] * n)
            tk.extend([k] * n)
            names.extend([name] * n)
            owners.extend([owner] * n)
            return ix
        if cfg.fixture == 'H2' or cfg.cells == 5:
            block('endplate', 0.01, cfg.end_n, 7900 * 500, 15, -1)
        pending_bp = None
        for cell in range(cfg.cells):
            left = pending_bp if pending_bp is not None else block('BP', 0.002, cfg.board_n, 1980 * 766, 95, cell)
            mea = []
            for i in range(5):
                mea.extend(block(LAYER_NAMES[i], THICKNESS[i], cfg.mesh[i], DRY_C[i], K_SOLID[i], cell))
            right = block('BP', 0.002, cfg.board_n, 1980 * 766, 95, cell)
            pending_bp = right if cfg.shared_bp and cell < cfg.cells - 1 else None
            self.mea_indices.append(mea)
            self.heater_indices.append(right)
            self.unit_indices.append(np.r_[left, mea, right])
        if cfg.fixture == 'H2' or cfg.cells == 5:
            block('endplate', 0.01, cfg.end_n, 7900 * 500, 15, -1)
        self.tdx, self.C, self.k0 = [np.array(a, dtype=float) for a in (td, tc, tk)]
        self.tx = np.cumsum(self.tdx) - self.tdx / 2
        self.names, self.owners = (np.array(names), np.array(owners))
        self.mea_indices = np.array(self.mea_indices)
        self.nt = len(td)

    def manifest(self):
        import pandas as pd
        return pd.DataFrame(dict(x_m=self.tx, dx_m=self.tdx, layer=self.names, cell=self.owners, dry_C_J_m3K=self.C, k_table=self.k0))

# ==================== model ====================
class Model:

    def __init__(self, cfg):
        cfg.validate()
        self.cfg = cfg
        self.g = Geometry(cfg)
        self.c, self.n = (cfg.cells, self.g.n)
        self.nw = 4 * self.c * self.n
        self.ih = self.nw
        self.il = self.ih + self.g.nt
        self.size = self.il + 3 * self.c + 2
        self.active_columns = []
        for block, mask in enumerate((self.g.porous, self.g.porous, self.g.omega > 0, self.g.porous)):
            for c in range(self.c):
                self.active_columns.extend(block * self.c * self.n + c * self.n + np.flatnonzero(mask))
        self.active_columns.extend(range(self.ih, self.il))
        self._jac_pattern = None
        self._jac_groups = None
        self._jac_scale = np.ones(self.size)
        self._jac_scale[:self.c * self.n] = 0.0001
        self._jac_scale[self.c * self.n:2 * self.c * self.n] = 0.01
        self._jac_scale[3 * self.c * self.n:self.nw] = 0.01

    def unpack(self, y):
        z = y[:self.nw].reshape(4, self.c, self.n)
        return (z, y[self.ih:self.il] * self.g.C, y[self.il:])

    def phases(self, y):
        z, H, led = self.unpack(y)
        w, ice, b, gas = z
        T = TM + H / self.g.C
        ix = self.g.mea_indices
        for _ in range(6):
            mv, ml, eg = partition(w, ice, T[ix], self.g.eps, self.cfg.vapor_equilibrium == 'supercooled_liquid')
            cap = self.g.C[ix] + 2000 * mv + 4182 * ml + 2050 * ice + 4182 * b
            new = TM + (H[ix] - LV * mv + LF * ice) / cap
            T[ix] = new
        mv, ml, eg = partition(w, ice, T[ix], self.g.eps, self.cfg.vapor_equilibrium == 'supercooled_liquid')
        lam = np.divide(b, self.g.bcap, out=np.zeros_like(b), where=self.g.bcap > 0)
        conc = gas / np.maximum(eg, 1e-12)
        return dict(T=T, local_T=T[ix], mv=mv, ml=ml, mi=ice, mb=b, eg=eg, lam=lam, gas=conc, led=led, H=H, mobile=w)

    def initial(self, Tfield=None):
        y = np.zeros(self.size)
        z = y[:self.nw].reshape(4, self.c, self.n)
        z[2] = self.g.bcap * self.cfg.lambda0
        T = np.full(self.g.nt, self.cfg.T0) if Tfield is None else np.array(Tfield).copy()
        assert T.shape == (self.g.nt,)
        z[3] = self.g.eps * self.g.gas_y * P / (R * T[self.g.mea_indices])
        H = self.g.C * (T - TM)
        H[self.g.mea_indices] += z[2] * 4182 * (T[self.g.mea_indices] - TM)
        y[self.ih:self.il] = H / self.g.C
        return y

    def thermal_k(self, s):
        k = self.g.k0.copy()
        if self.cfg.thermal == 'effective':
            return k
        layer = self.g.layer
        ks = np.take(K_SOLID, layer)
        kg = np.where(layer < 2, 0.1672, YO2 * 0.0246 + (1 - YO2) * 0.0235)
        solid = 1 - self.g.eps - self.g.omega
        km = solid * ks + self.g.omega * 0.24 + s['eg'] * kg + s['ml'] / 990 * 0.6 + s['mi'] / 920 * 2.3
        km[:, self.g.pem] = 0.24
        k[self.g.mea_indices] = km
        return k

    def voltage(self, s, j):
        g = self.g
        T = s['local_T']
        J = j * 10000.0
        tc = np.average(T[:, g.clc], axis=1, weights=g.dx[g.clc])
        ca = np.average(s['gas'][:, g.cla], axis=1, weights=g.dx[g.cla])
        co = np.average(s['gas'][:, g.clc], axis=1, weights=g.dx[g.clc])
        ta = np.average(T[:, g.cla], axis=1, weights=g.dx[g.cla])
        ph, po = (ca * R * ta / P, co * R * tc / P)
        E = 1.229 - 0.00085 * (tc - 298.15) + R * tc / (2 * F) * np.log(np.maximum(ph, 1e-20) * np.sqrt(np.maximum(po, 1e-20)))
        si = s['mi'][:, g.clc] / (920 * g.eps[g.clc])
        hydration = np.clip(s['lam'][:, g.clc] / 14.0, 1e-12, 1.0) ** self.cfg.cathode_hydration_exponent
        active = np.average(np.maximum(1 - si, 1e-12) ** 3.5 * hydration, axis=1, weights=g.dx[g.clc])
        exchange = self.cfg.j0 * np.exp(-67000 / R * (1 / tc - 1 / 298.15)) * active
        act = R * tc / (0.5 * F) * np.arcsinh(J / (2 * exchange))
        kap = conductivity(T[:, g.pem], s['lam'][:, g.pem])
        ohm = J * (np.sum(g.dx[g.pem] / np.maximum(kap, 1e-12), axis=1) + self.cfg.Rc)
        cl_ohm = np.zeros(self.c)
        cl_kmin = np.full(self.c, np.inf)
        if self.cfg.cl_proton_loss:
            ids = np.r_[g.cla, g.clc]
            kcl = conductivity(T[:, ids], s['lam'][:, ids]) * g.omega[ids] ** 1.5
            lengths = np.where(g.layer[ids] == 1, THICKNESS[1], THICKNESS[3])
            weights = g.im[ids] ** 2 + (g.dx[ids] / lengths) ** 2 / 12
            cl_ohm = J * np.sum(g.dx[ids] * weights / np.maximum(kcl, 1e-12), axis=1)
            cl_kmin = np.min(kcl, axis=1)
            ohm += cl_ohm
        Dc = 2.2e-05 * (T[:, g.cathode] / 298.15) ** 1.75 * np.maximum(s['eg'][:, g.cathode], 1e-12) ** 1.5
        path = THICKNESS[4] + THICKNESS[3] * (0.5 if self.cfg.half_cl else 1.0)
        deff = np.sum(g.dx[g.cathode]) / np.sum(g.dx[g.cathode] / Dc, axis=1)
        jlim = 4 * F * deff * co / path
        ratio = J / np.maximum(jlim, 1e-20)
        con = -R * tc / (4 * F) * np.log(np.maximum(1 - ratio, 1e-12))
        if self.cfg.concentration == 'unified':
            con[:] = 0
        if self.c == 5:
            con[[0, -1]] *= self.cfg.end_concentration_factor
        V = E - act - ohm - con
        return dict(V=V, reversible=E, activation=act, ohmic=ohm, concentration=con, cl_ohmic=cl_ohm, active_fraction=active, jlim=jlim / 10000.0, ratio=ratio, kappa_min=np.minimum(np.min(kap, axis=1), cl_kmin))

    def observations(self, s, j):
        g = self.g
        tm = np.average(s['local_T'], axis=1, weights=g.dx)
        tu = np.array([np.average(s['T'][ix], weights=g.tdx[ix]) for ix in g.unit_indices])
        temp = tm if self.cfg.observation == 'MEA' else tu
        return dict(T_mean=temp, T_mean_MEA=tm, T_mean_unit=tu, ice_bulk=np.max(s['mi'] / 920, axis=1), ice_saturation=np.max(s['mi'][:, g.porous] / (920 * g.eps[g.porous]), axis=1), **self.voltage(s, j))

    def margins(self, y, j):
        s = self.phases(y)
        o = self.observations(s, j)
        g = self.g
        D = membrane_D(s['local_T'][:, g.bound], s['lam'][:, g.bound])
        return dict(voltage=float(np.min(o['V']) - self.cfg.voltage_limit), pore=float(np.min(s['eg'][:, g.porous]) - 1e-08), gas=float(np.min(s['gas'][:, g.porous]) - 1e-08), oxygen=float(1 - np.max(o['ratio'])), membrane=float(min(np.min(o['kappa_min']), np.min(D) * 10000000000.0)), nonnegative=float(min(np.min(s['mobile']), np.min(s['mi']), np.min(s['mb'])) + 1e-05), charge=float(self.cfg.charge_limit - s['led'][-1]), success=float(np.min(o['T_mean']) - TM - self.cfg.success_margin))

    def rhs(self, t, y, j, u, details=False):
        g = self.g
        cfg = self.cfg
        J = j * 10000.0
        s = self.phases(y)
        v = self.voltage(s, j)
        out = np.zeros_like(y)
        dz = out[:self.nw].reshape(4, self.c, self.n)
        dw, di, db, dg = dz
        dH = np.zeros(g.nt)
        ledger = out[self.il:]
        massheat = np.zeros((self.c, self.n))
        fluxes = {name: np.zeros((self.c, self.n + 1)) for name in ('vapor', 'liquid', 'bound', 'gas')}
        water_external = np.zeros(self.c)
        net_energy_mass = 0.0
        for c in range(self.c):
            T = s['local_T'][c]
            hv, hl, hi, hb = water_enthalpies(T)
            for ids, side in ((g.anode, 0), (g.cathode, -1)):
                dx = g.dx[ids]
                eg = np.maximum(s['eg'][c, ids], 1e-12)
                gasD = g.gas_ref[ids] * (T[ids] / 298.15) ** 1.75 * eg ** 1.5
                cv = s['mv'][c, ids] / (MW * eg)
                dv = g.vap_ref[ids] * (T[ids] / 298.15) ** 1.75 * eg ** 1.5
                fv = np.zeros(len(ids) + 1)
                fg = fv.copy()
                fv[1:-1] = MW * harmonic_faces(dx, dv) * (cv[:-1] - cv[1:])
                cg = s['gas'][c, ids]
                fg[1:-1] = harmonic_faces(dx, gasD) * (cg[:-1] - cg[1:])
                if side == 0:
                    fv[0] = -MW * dv[0] * cv[0] / (dx[0] / 2)
                    fg[0] = gasD[0] * (g.gas_y[ids[0]] * P / (R * T[ids[0]]) - cg[0]) / (dx[0] / 2)
                else:
                    fv[-1] = MW * dv[-1] * cv[-1] / (dx[-1] / 2)
                    fg[-1] = gasD[-1] * (cg[-1] - g.gas_y[ids[-1]] * P / (R * T[ids[-1]])) / (dx[-1] / 2)
                fl = np.zeros_like(fv)
                if cfg.water == 'W1':
                    fl = liquid_flux(s['ml'][c, ids], s['mi'][c, ids], T[ids], g.eps[ids], g.K[ids], g.angle[ids], dx, cfg.mu_factor, cfg.ice_connectivity, cfg.outlet == 'drain', side)
                dw[c, ids] += (fv[:-1] + fl[:-1] - fv[1:] - fl[1:]) / dx
                dg[c, ids] += (fg[:-1] - fg[1:]) / dx
                for name, f in [('vapor', fv), ('liquid', fl), ('gas', fg)]:
                    fluxes[name][c, ids[0]:ids[-1] + 2] = f
                water_external[c] += fv[0] + fl[0] - fv[-1] - fl[-1]
                e = np.zeros_like(fv)
                for f, h in [(fv, hv[ids]), (fl, hl[ids])]:
                    e[1:-1] += f[1:-1] * np.where(f[1:-1] >= 0, h[:-1], h[1:])
                    e[0] += f[0] * h[0]
                    e[-1] += f[-1] * h[-1]
                massheat[c, ids] += (e[:-1] - e[1:]) / dx
                net_energy_mass += e[0] - e[-1]
            ids = g.bound
            lam = s['lam'][c, ids]
            D = np.maximum(membrane_D(T[ids], lam), 1e-20)
            conduct = g.omega[ids] ** 1.5 * 2150 * MW * D
            f = np.zeros(len(ids) + 1)
            f[1:-1] = harmonic_faces(g.dx[ids], conduct) * (lam[:-1] - lam[1:])
            fm = (g.im[ids[:-1]] + g.im[ids[1:]]) / 2
            donor = lam[:-1]
            f[1:-1] += MW * (2.5 * donor / 22) * J * fm / F
            db[c, ids] += (f[:-1] - f[1:]) / g.dx[ids]
            fluxes['bound'][c, ids[0]:ids[-1] + 2] = f
            e = np.zeros_like(f)
            e[1:-1] = f[1:-1] * np.where(f[1:-1] >= 0, hb[ids[:-1]], hb[ids[1:]])
            massheat[c, ids] += (e[:-1] - e[1:]) / g.dx[ids]
            cls = np.r_[g.cla, g.clc]
            pv = s['mv'][c, cls] / np.maximum(s['eg'][c, cls], 1e-12) * R * T[cls] / MW
            beq = g.bcap[cls] * lambda_eq(pv / psat(T[cls], liquid=True))
            rb = (beq - s['mb'][c, cls]) / cfg.tau_b
            dw[c, cls] -= rb
            db[c, cls] += rb
            freeze = s['ml'][c] * np.maximum(TM - T, 0) / (20 * cfg.tau_f)
            melt = np.maximum(s['mi'][c], 0) * np.maximum(T - TM, 0) / (5 * cfg.tau_m)
            di[c] = freeze - melt
            dw[c] -= di[c]
            sw = MW * J / (2 * F * THICKNESS[3])
            dw[c, g.clc] += sw
            massheat[c, g.clc] += sw * hl[g.clc]
            net_energy_mass += np.sum(sw * hl[g.clc] * g.dx[g.clc])
            dg[c, g.cla] -= J / (2 * F * THICKNESS[1])
            dg[c, g.clc] -= J / (4 * F * THICKNESS[3])
            qrx = J * (1.48 - v['V'][c])
            dH[g.mea_indices[c]] += qrx / sum(THICKNESS) + massheat[c]
            heater = g.heater_indices[c]
            dH[heater] += u[c] * 10000.0 / sum(g.tdx[heater])
            ledger[c] = MW * J / (2 * F) + water_external[c]
            ledger[self.c + c] = qrx
            ledger[2 * self.c + c] = u[c] * 10000.0
        k = self.thermal_k(s)
        conduct = harmonic_faces(g.tdx, k)
        if cfg.contact_r:
            interface = (g.names[:-1] == 'endplate') != (g.names[1:] == 'endplate')
            conduct[interface] = 1 / (1 / conduct[interface] + cfg.contact_r)
        heatflux = np.zeros(g.nt + 1)
        heatflux[1:-1] = conduct * (s['T'][:-1] - s['T'][1:])
        heatflux[0] = (cfg.ambient - s['T'][0]) / (1 / cfg.h + g.tdx[0] / (2 * k[0])) if cfg.h else 0.0
        heatflux[-1] = (s['T'][-1] - cfg.ambient) / (1 / cfg.h + g.tdx[-1] / (2 * k[-1])) if cfg.h else 0.0
        dH += (heatflux[:-1] - heatflux[1:]) / g.tdx
        out[self.ih:self.il] = dH / g.C
        ledger[-2] = np.sum(ledger[self.c:3 * self.c]) + net_energy_mass + heatflux[0] - heatflux[-1]
        ledger[-1] = j
        if details:
            return dict(derivative=out, fluxes=fluxes, heatflux=heatflux, water_external=water_external, reaction_W_m2=ledger[self.c:2 * self.c], massheat=massheat)
        return out

    def balances(self, y, initial):
        z, H, led = self.unpack(y)
        zi, Hi, _ = self.unpack(initial)
        dw = np.sum((z[:3] - zi[:3]) * self.g.dx, axis=(0, 2))
        re_w = dw - led[:self.c]
        re_h = np.dot(H - Hi, self.g.tdx) - led[-2]
        wscale = np.maximum(np.sum(zi[:3] * self.g.dx, axis=(0, 2)) + MW * 10000.0 * led[-1] / (2 * F), 1e-08)
        hscale = max(abs(led[-2]), abs(np.dot(Hi, self.g.tdx)), 1.0)
        return (float(np.max(abs(re_w) / wscale)), float(abs(re_h) / hscale))

    def sparsity(self):
        a = lil_matrix((self.size, self.size), dtype=int)
        g = self.g
        for c in range(self.c):
            ix = np.r_[np.concatenate([np.arange(b * self.c * self.n + c * self.n, b * self.c * self.n + (c + 1) * self.n) for b in range(4)]), self.ih + g.mea_indices[c]]
            a[np.ix_(ix, ix)] = 1
        for i in range(g.nt):
            a[self.ih + i, self.ih + max(0, i - 1):self.ih + min(g.nt, i + 2)] = 1
        a[self.il:, :self.il] = 1
        return a.tocsr()

    def jacobian(self, t, y, j, u):
        """Bounded forward differences avoid unbounded perturbations of zero slots.

        All ledger columns are analytically zero. The invariant off-material slots
        remain zero, so their Newton corrections and columns can also be omitted.
        A fixed relative step is independent of SciPy's adaptive num_jac factors.
        """
        if self._jac_pattern is None:
            pattern = lil_matrix((self.size, self.size), dtype=int)
            for c in range(self.c):
                ix = np.r_[np.concatenate([np.arange(b * self.c * self.n + c * self.n, b * self.c * self.n + (c + 1) * self.n) for b in range(4)]), self.ih + self.g.unit_indices[c]]
                pattern[np.ix_(ix, ix)] = 1
                pattern[self.il + c, ix] = 1
                pattern[self.il + self.c + c, ix] = 1
            for i in range(self.g.nt):
                pattern[self.ih + i, self.ih + max(0, i - 1):self.ih + min(self.g.nt, i + 2)] = 1
            self._jac_pattern = pattern.tocsc()
            occupied = []
            groups = []
            for col in self.active_columns:
                ix = set(self._jac_pattern.indices[self._jac_pattern.indptr[col]:self._jac_pattern.indptr[col + 1]])
                for k, used in enumerate(occupied):
                    if not ix & used:
                        groups[k].append(col)
                        used.update(ix)
                        break
                else:
                    groups.append([col])
                    occupied.append(ix)
            self._jac_groups = groups
        f0 = self.rhs(t, y, j, u)
        rows = []
        cols = []
        data = []
        for group in self._jac_groups:
            yp = y.copy()
            delta = 1e-08 * np.maximum(self._jac_scale[group], np.abs(y[group]))
            yp[group] += delta
            diff = self.rhs(t, yp, j, u) - f0
            for col, step in zip(group, delta):
                ix = self._jac_pattern.indices[self._jac_pattern.indptr[col]:self._jac_pattern.indptr[col + 1]]
                values = diff[ix] / step
                nz = values != 0
                ix = ix[nz]
                values = values[nz]
                rows.extend(ix)
                cols.extend([col] * len(ix))
                data.extend(values)
        matrix = csc_matrix((data, (rows, cols)), shape=(self.size, self.size))
        heat_row = np.asarray(self.g.C * self.g.tdx @ matrix[self.ih:self.il, :]).ravel()
        nz = np.flatnonzero(heat_row)
        matrix += csc_matrix((heat_row[nz], (np.full(len(nz), self.size - 2), nz)), shape=matrix.shape)
        return matrix

# ==================== protocols ====================
@dataclass
class Protocol:
    kind: str = 'constant'
    params: tuple = (0.1,)
    powers: tuple = (0.0,) * 5
    heat_off: float = 0.0
    load_start: float = 0.0
    times: object = None
    currents: object = None

    def j(self, t):
        if self.kind == 'experiment':
            return float(np.interp(t, self.times, self.currents))
        s = max(0.0, t - self.load_start)
        if t < self.load_start:
            return 0.0
        if self.kind == 'constant':
            return self.params[0]
        if self.kind == 'ramp':
            a, b, tr = self.params
            return a + (b - a) * min(s / max(tr, 1e-08), 1)
        if self.kind == 'step':
            a, b, c, t1, t2 = self.params
            return a if s < t1 else b if s < t2 else c
        if self.kind == 'fixed':
            return min(0.005 * s, 0.3)
        if self.kind == 'zero':
            return 0.0
        raise ValueError(self.kind)

    def u(self, t, cells):
        return np.asarray(self.powers[:cells], float) if t < self.heat_off else np.zeros(cells)

    def breaks(self, end):
        b = [0.0, end, self.heat_off, self.load_start]
        if self.kind == 'step':
            b += [self.load_start + self.params[3], self.load_start + self.params[4]]
        if self.kind in ('ramp', 'fixed'):
            b += [self.load_start + (60 if self.kind == 'fixed' else self.params[2])]
        return sorted(set((float(v) for v in b if 0 <= v <= end)))

    def valid(self, cfg):
        if len(self.powers) < cfg.cells or not np.isfinite(self.powers).all():
            return False
        if min(self.powers) < 0 or max(self.powers) > 1 or self.heat_off < 0:
            return False
        if self.kind == 'experiment':
            levels = np.asarray(self.currents)
        elif self.kind == 'constant':
            levels = np.asarray(self.params[:1])
        elif self.kind == 'ramp':
            levels = np.asarray(self.params[:2])
        elif self.kind == 'step':
            levels = np.asarray(self.params[:3])
        elif self.kind == 'fixed':
            levels = np.array([0.0, 0.3])
        elif self.kind == 'zero':
            levels = np.array([0.0])
        else:
            return False
        return bool(np.isfinite(levels).all() and np.min(levels) >= 0 and (np.max(levels) <= cfg.current_limit + 1e-12))

# ==================== solver ====================
def simulate(cfg, protocol, Tfield=None, controller=None, stop_success=True, initial=None, complete_horizon=False):
    model = Model(cfg)
    y0 = model.initial(Tfield) if initial is None else initial.copy()
    y = y0.copy()
    ts = None
    reason = 'horizon'
    wall = time.perf_counter()
    times, states, currents, powers = ([], [], [], [])
    bounds = protocol.breaks(cfg.horizon)

    def append(t, state, j, u):
        times.append(float(t))
        states.append(state.copy())
        currents.append(j)
        powers.append(np.array(u).copy())

    def events_for(jfun):
        keys = ['voltage', 'pore', 'gas', 'oxygen', 'membrane', 'nonnegative', 'charge', 'success']
        events = []
        for key in keys:

            def event(t, state, key=key):
                return model.margins(state, jfun(t))[key] + (1e-09 if key == 'voltage' else 0.0)
            event.direction = 1 if key == 'success' else -1
            heat_finished = controller.off_time is not None if controller is not None else a >= protocol.heat_off
            event.terminal = key != 'success' or (not complete_horizon and (stop_success or heat_finished))
            events.append(event)
        return (keys, events)
    nfev = 0
    extrema = dict(V_min=float('inf'), ice_max=0.0, delta_T_max=0.0)
    extrema_to_success = None

    def update_ext(t, state, j):
        nonlocal extrema_to_success
        obs = model.observations(model.phases(state), j)
        extrema['V_min'] = min(extrema['V_min'], float(min(obs['V'])))
        extrema['ice_max'] = max(extrema['ice_max'], float(max(obs['ice_bulk'])))
        extrema['delta_T_max'] = max(extrema['delta_T_max'], float(np.ptp(obs['T_mean'])))
        if ts is None or t <= ts + 1e-08:
            extrema_to_success = extrema.copy()
    try:
        a = 0.0
        while a < cfg.horizon - 1e-10:
            if not protocol.valid(cfg):
                reason = 'control_input'
                break
            b = next((x for x in bounds if x > a + 1e-10))
            if controller is not None:
                b = min(b, a + cfg.control_prediction_horizon)
            j0 = protocol.j(a)
            if controller is not None:
                obs = model.observations(model.phases(y), j0)
                if controller.last_t is not None and abs(controller.last_t - a) < 1e-09:
                    u = controller.previous_power.copy()
                else:
                    u = controller.step(a, obs['T_mean'] - TM, obs['V'], j0)
            else:
                u = protocol.u((a + b) / 2, cfg.cells)
            append(a, y, j0, u)
            update_ext(a, y, j0)
            margins = model.margins(y, j0)
            failed = [k for k, value in margins.items() if k != 'success' and value < -1e-09]
            if failed:
                reason = failed[0]
                break
            if ts is None and margins['success'] >= 0:
                ts = a
                heat_finished = controller.off_time is not None if controller is not None else a >= protocol.heat_off
                if not complete_horizon and (stop_success or heat_finished):
                    reason = 'success'
                    break

            def jf(t):
                return protocol.j(min(t, np.nextafter(b, a)))
            keys, events = events_for(jf)
            atol = np.full(model.size, 1e-07)
            atol[2 * model.c * model.n:3 * model.c * model.n] = 1e-06
            atol[model.ih:model.il] = 1e-07
            atol[model.il:] = 1e-09
            sol = solve_ivp(lambda t, state: model.rhs(t, state, jf(t), u), (a, b), y, method='BDF', rtol=cfg.rtol, atol=atol, max_step=cfg.max_step, jac=lambda t, state: model.jacobian(t, state, jf(t), u), dense_output=True, events=events)
            nfev += sol.nfev
            accepted_end = float(sol.t[-1])
            if controller is not None:
                first = (np.floor((a + 1e-09) / cfg.sample) + 1) * cfg.sample
                for tc in np.arange(first, sol.t[-1] + 1e-09, cfg.sample):
                    tc = min(float(tc), float(sol.t[-1]))
                    sample_state = sol.sol(tc)
                    obs = model.observations(model.phases(sample_state), protocol.j(tc))
                    new_u = controller.step(tc, obs['T_mean'] - TM, obs['V'], protocol.j(tc))
                    if not np.array_equal(new_u, u):
                        accepted_end = tc
                        break
            for ie, key in enumerate(keys):
                if len(sol.t_events[ie]) and sol.t_events[ie][0] <= accepted_end + 1e-09:
                    te = float(sol.t_events[ie][0])
                    if key == 'success' and ts is None:
                        ts = te
                    if events[ie].terminal:
                        reason = key
            check_times = list(sol.t[sol.t <= accepted_end + 1e-09]) + [accepted_end]
            if ts is not None and a <= ts <= accepted_end:
                check_times.append(ts)
            for t in sorted(set(check_times)):
                update_ext(t, sol.sol(t), jf(t))
            tt = np.arange(a + cfg.sample, accepted_end - 1e-09, cfg.sample)
            output_times = list(tt) + [accepted_end]
            if ts is not None and a <= ts <= accepted_end:
                output_times.append(ts)
            for t in sorted(set(output_times)):
                append(t, sol.sol(t), jf(t), u)
            y = sol.sol(accepted_end)
            if not sol.success:
                reason = 'numerical_failure'
                break
            if sol.status == 1 and accepted_end >= sol.t[-1] - 1e-09:
                marg = model.margins(y, jf(sol.t[-1]))
                if reason == 'charge' and marg['success'] >= -1e-07 and (marg['voltage'] >= -1e-08):
                    ts = float(sol.t[-1])
                    reason = 'success'
                break
            if not complete_horizon and controller is not None and (ts is not None) and (controller.off_time is not None) and (accepted_end >= max(ts, controller.off_time)):
                reason = 'success'
                break
            if not complete_horizon and (not stop_success) and (ts is not None) and (controller is None) and (accepted_end >= max(ts, protocol.heat_off)):
                reason = 'success'
                break
            a = accepted_end
    except (ValueError, RuntimeError, FloatingPointError, OverflowError, np.linalg.LinAlgError) as exc:
        reason = 'numerical_failure'
        exception = repr(exc)
    else:
        exception = None
    if not times:
        append(0.0, y, protocol.j(0), protocol.u(0, cfg.cells))
    T, V, ice, si, lam, ph = ([], [], [], [], [], [])
    balances = []
    for state, j in zip(states, currents):
        s = model.phases(state)
        obs = model.observations(s, j)
        T.append(obs['T_mean'] - TM)
        V.append(obs['V'])
        ice.append(obs['ice_bulk'])
        si.append(obs['ice_saturation'])
        lam.append(np.average(s['lam'][:, model.g.pem], axis=1, weights=model.g.dx[model.g.pem]))
        balances.append(model.balances(state, y0))
    T, V, ice, si, lam, balances = map(np.array, (T, V, ice, si, lam, balances))
    ledger = y[model.il:]
    last = times[-1]
    off = controller.off_time if controller is not None else min(protocol.heat_off, last) if protocol.heat_off <= last else None
    if protocol.kind == 'zero' and stop_success and (ts is not None):
        off = ts
    energies = ledger[2 * model.c:3 * model.c] * AREA
    safe_reason = reason in ('success', 'horizon')
    feasible = ts is not None and safe_reason and (off is not None)
    feasibility = 'feasible' if feasible else 'unresolved' if reason in ('horizon', 'numerical_failure') else 'infeasible'
    official = ts is not None
    result = dict(case_hash=cfg.digest, feasibility=feasibility, stop_reason=reason, success_time=ts, last_valid_time=last, th_startup=off, startup_result='success' if official else reason, shutdown_status='passed' if feasible else 'unresolved' if safe_reason else reason, E_aux_total_J=float(sum(energies)), E_cells_J=energies.tolist(), charge_C_cm2=float(ledger[-1]), wall_s=time.perf_counter() - wall, nfev=nfev, mass_residual=float(max(balances[:, 0])), energy_residual=float(max(balances[:, 1])), exception=exception, **extrema_to_success or extrema)
    if ts is not None:
        k = np.searchsorted(np.asarray(times), ts)
        k = min(k, len(states) - 1)
        result['E_aux_to_success_J'] = float(sum(states[k][model.il + 2 * model.c:model.il + 3 * model.c]) * AREA)
    result['shutdown_extrema'] = extrema
    return dict(metrics=result, config=cfg.to_dict(), t=np.array(times), y=np.array(states), j=np.array(currents), u=np.array(powers), T=T, V=V, ice=ice, ice_saturation=si, lambda_mean=lam, balances=balances, x=model.g.tx, layer=model.g.names, cell=model.g.owners)

# ==================== io ====================
io_ROOT = Path(os.environ.get('B_MODEL_WORKSPACE', str(Path(__file__).resolve().parent))).resolve()

PLAN_ROOT = io_ROOT.parent

io_SOURCE = PLAN_ROOT.parent / '中文题目' / 'B题.zip'

if not io_SOURCE.exists():
    io_SOURCE = io_ROOT / 'data' / 'original' / 'B题.zip'

OUTPUT = _project_storage() / 'outputs'

def json_default(v):
    if isinstance(v, np.ndarray):
        return v.tolist()
    if isinstance(v, np.generic):
        return v.item()
    if isinstance(v, Path):
        return str(v)
    raise TypeError(type(v).__name__)

def write_json(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, default=json_default), encoding='utf-8')

def load_source(path=io_SOURCE):
    blob = Path(path).read_bytes()
    digest = hashlib.sha256(blob).hexdigest()
    expected = '982dd5c1f141a1c76438e8b165c3d049c82ac9ef4b4b416eb740181ca592447f'
    if digest != expected:
        raise ValueError('Source archive changed; rerun provenance review')
    tables = []
    params = []
    experiments = {}
    formulas = []
    with zipfile.ZipFile(io.BytesIO(blob)) as outer:
        for name in outer.namelist():
            if name.endswith('.docx'):
                with zipfile.ZipFile(io.BytesIO(outer.read(name))) as doc:
                    root = ET.fromstring(doc.read('word/document.xml'))
                    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
                    for table in root.findall('.//w:tbl', ns):
                        tables.append([[''.join(el.itertext()) for el in cell.findall('.//w:t', ns)] for cell in []])
                        tables[-1] = [[''.join((t.text or '' for t in cell.findall('.//w:t', ns))) for cell in row.findall('w:tc', ns)] for row in table.findall('w:tr', ns)]
            if name.endswith('.zip'):
                with zipfile.ZipFile(io.BytesIO(outer.read(name))) as inner:
                    for fn in inner.namelist():
                        if not fn.endswith('.xlsx'):
                            continue
                        data = inner.read(fn)
                        book = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
                        form = openpyxl.load_workbook(io.BytesIO(data), data_only=False)
                        for sheet in book:
                            if sheet.max_row < 100:
                                for row in range(2, sheet.max_row + 1):
                                    params.append(dict(source_workbook=fn, source_cell=f'{sheet.title}!C{row}', category=sheet.cell(row, 1).value, name=sheet.cell(row, 2).value, raw_value=sheet.cell(row, 3).value, unit=sheet.cell(row, 4).value, note=sheet.cell(row, 5).value))
                            else:
                                rows = []
                                for row in range(3, sheet.max_row + 1):
                                    vals = [sheet.cell(row, k).value for k in range(1, 6)]
                                    rows.append([fn, sheet.title, row, *vals])
                                    formulas.append(dict(sheet=sheet.title, row=row, formula=form[sheet.title].cell(row, 5).value))
                                frame = pd.DataFrame(rows, columns=['source_workbook', 'source_sheet', 'source_row', 't', 'I', 'V', 'T', 'j'])
                                frame['j_A'] = frame.I / 25
                                frame['I_B'] = 25 * frame.j
                                assert len(frame) == 184 and np.isfinite(frame[['t', 'I', 'V', 'T', 'j']]).all().all()
                                assert np.all(np.diff(frame.t) > 0)
                                temp = -25 if '25' in sheet.title else -20
                                experiments[temp] = frame
    return dict(sha256=digest, parameters=params, experiments=experiments, tables=tables, formulas=formulas)

def audit():
    src = load_source()
    out = io_ROOT / 'data' / 'processed'
    out.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(src['parameters']).to_csv(out / 'parameter_source.csv', index=False, encoding='utf-8-sig')
    for temp, d in src['experiments'].items():
        d.to_csv(out / f'experiment_{temp}.csv', index=False, encoding='utf-8-sig')
    write_json(out / 'source_provenance.json', {k: v for k, v in src.items() if k != 'experiments'})
    return src

def save_run(result, name, protocol=None):
    path = _project_storage() / 'runs' / name
    path.mkdir(parents=True, exist_ok=True)
    write_json(path / 'metrics.json', result['metrics'])
    write_json(path / 'config.json', result['config'])
    if protocol is not None:
        write_json(path / 'protocol.json', vars(protocol))
    np.savez_compressed(path / 'trajectory.npz', **{k: v for k, v in result.items() if k not in ('metrics', 'config')})
    rows = []
    for i, t in enumerate(result['t']):
        for c in range(result['T'].shape[1]):
            rows.append(dict(t_s=t, cell=c + 1, j_A_cm2=result['j'][i], u_W_cm2=result['u'][i, c], T_C=result['T'][i, c], V_V=result['V'][i, c], ice_bulk=result['ice'][i, c], ice_saturation=result['ice_saturation'][i, c], lambda_mean=result['lambda_mean'][i, c], water_residual=result['balances'][i, 0], energy_residual=result['balances'][i, 1]))
    pd.DataFrame(rows).to_csv(path / 'timeseries.csv', index=False)
    return path

def load_run(name):
    path = _project_storage() / 'runs' / name
    result = dict(np.load(path / 'trajectory.npz', allow_pickle=False))
    result['metrics'] = json.loads((path / 'metrics.json').read_text(encoding='utf-8'))
    result['config'] = json.loads((path / 'config.json').read_text(encoding='utf-8'))
    return result

def environment():
    from importlib.metadata import distributions
    packages = sorted((f"{d.metadata['Name']}=={d.version}" for d in distributions()))
    (io_ROOT / 'environment.lock').write_text('\n'.join((p for p in packages if not p.startswith('fcstart'))) + '\n', encoding='utf-8')
    write_json(OUTPUT / 'environment.json', dict(python=sys.version, platform=platform.platform(), packages=packages))

# ==================== evaluate ====================
def _snapshot_code():
    h = hashlib.sha256()
    for name in ('q1.py', 'q2.py', 'q3.py', 'q4.py'):
        p = Path(__file__).resolve().parent / name
        h.update(name.encode())
        h.update(p.read_bytes())
    return h.hexdigest()[:16]

_LOADED_CODE_DIGEST = _snapshot_code()

def code_digest():
    return _LOADED_CODE_DIGEST

def evaluate(cfg, protocol, Tfield=None, controller=None, stop_success=True, label='candidate', cache=True):
    key = dict(config=cfg.to_dict(), protocol=vars(protocol), Tfield=Tfield, controller=None if controller is None else controller.metadata(), code=code_digest(), stop_success=stop_success)
    digest = hashlib.sha256(json.dumps(key, sort_keys=True, default=json_default).encode()).hexdigest()[:20]
    name = f'{label}_{digest}'
    path = _project_storage() / 'runs' / name
    hit = cache and (path / 'metrics.json').exists() and (controller is None)
    if hit:
        result = load_run(name)
    else:
        retries = []
        for attempt in range(3):
            used = cfg.changed(max_step=cfg.max_step / 2 ** attempt, rtol=cfg.rtol / 2 ** attempt)
            result = simulate(used, protocol, Tfield, controller, stop_success)
            retries.append(dict(attempt=attempt, reason=result['metrics']['stop_reason'], wall_s=result['metrics']['wall_s']))
            if result['metrics']['stop_reason'] != 'numerical_failure' or controller is not None:
                break
        result['metrics']['retries'] = retries
        result['metrics']['candidate_id'] = name
        result['metrics']['numerically_accepted'] = result['metrics']['mass_residual'] < 0.001 and result['metrics']['energy_residual'] < 0.001
        if not result['metrics']['numerically_accepted']:
            result['metrics']['feasibility'] = 'unresolved'
        result['metrics']['config_requested'] = cfg.digest
        result['metrics']['code_hash'] = key['code']
        save_run(result, name, protocol)
        if controller is not None:
            write_json(path / 'controller.json', controller.metadata())
            write_json(path / 'controller_log.json', controller.logs)
    write_json(_project_storage() / 'evaluation_ledger' / f'{time.time_ns()}_{digest}.json', dict(candidate=name, cache_hit=hit, **result['metrics']))
    return result

def score(result, objective='time'):
    r = result['metrics']
    if r['feasibility'] == 'feasible':
        return r['success_time'] if objective == 'time' else r['E_aux_total_J']
    if r['stop_reason'] == 'numerical_failure' or not r.get('numerically_accepted', True):
        return float('inf')
    temp = float(np.min(result['T'][-1]))
    return 1000000.0 + 10000.0 * max(-temp, 0) + 100000.0 * max(0.3 - r['V_min'], 0) + 1000 / (r['last_valid_time'] + 1)

# ==================== plot_style ====================
def configure_chinese_plotting(size=8):
    available = {font.name for font in font_manager.fontManager.ttflist}
    candidates = ('Microsoft YaHei', 'Noto Sans CJK SC', 'Source Han Sans SC', 'SimHei')
    family = next((name for name in candidates if name in available), None)
    if family is None:
        raise RuntimeError('请安装微软雅黑、思源黑体或 Noto Sans CJK SC 后重新绘图。')
    rcParams.update({'font.family': 'sans-serif', 'font.sans-serif': [family, 'DejaVu Sans'], 'font.size': size, 'axes.spines.top': False, 'axes.spines.right': False, 'legend.frameon': False, 'svg.fonttype': 'none', 'pdf.fonttype': 42, 'axes.linewidth': 0.8, 'axes.unicode_minus': True})
    return font_manager.findfont(font_manager.FontProperties(family=family))

STRATEGIES = {'constant': '恒流加载', 'ramp': '线性升载', 'step': '阶梯加载', 'preheat': '纯预加热', 'coheat': '恒功率协同加热', 'dynamic_speed_tradeoff': '动态功率控制（启动时间折中）'}

SCENARIOS = {'equilibrium': '完全冷却', '20min': '预冷20 min', '40min': '预冷40 min'}

STATUSES = {'feasible': '可行', 'infeasible': '不可行', 'unresolved': '未判定'}

STOP_REASONS = {'voltage': '电压下限', 'charge': '电荷上限', 'horizon': '计算时限', 'membrane': '膜模型适用域边界', 'success': '启动成功', 'pore': '孔隙下限', 'gas': '气体浓度下限', 'oxygen': '氧传输极限', 'nonnegative': '状态非负约束', 'numerical_failure': '数值失败', 'control_input': '控制输入违规'}

SENSITIVITY = {'baseline': '基准模型', 'tau_b_0.1x': '吸附时间常数×0.1', 'tau_b_10x': '吸附时间常数×10', 'tau_f_0.1x': '冻结时间常数×0.1', 'tau_f_10x': '冻结时间常数×10', 'mu_2x': '液水黏度×2', 'mu_5x': '液水黏度×5', 'ice_connectivity_1': '冰连通指数为1', 'ice_connectivity_5': '冰连通指数为5', 'effective_k': '等效导热系数', 'unit_temperature': '重复单元平均温度', 'unified_concentration': '统一浓差处理', 'half_CL_path': '半催化层扩散路径', 'h_20': '换热系数20 W/(m²·K)', 'h_60': '换热系数60 W/(m²·K)'}

FIGURE_NAMES = {'grid_convergence': '网格收敛性', 'parameter_sensitivity': '参数敏感性', 'Q1_calibration': '问题1_模型校准与实验对比', 'Q1_inferred_states': '问题1_模型推断状态', 'Q1_temperature_holdout': '问题1_温度留出检验', 'Q3_energy_time_tradeoff': '问题3_条件性能耗与时间权衡', 'Q4_precool_startup_scan': '问题4_预冷时长与启动结果', 'Q4_precooling_spatial': '问题4_预冷空间温度分布', 'Q4_precooling': '问题4_预冷温度变化', '电压预测优化对比': '电压预测优化对比'}

SEARCH_NAMES = {}

for question, kinds in ((2, ('constant', 'ramp', 'step')), (3, ('preheat', 'coheat'))):
    for kind in kinds:
        prefix = f'q{question}_{kind}'
        SEARCH_NAMES[prefix] = f'问题{question}：{STRATEGIES[kind]}搜索'
        FIGURE_NAMES[prefix + '_evaluations_search_history'] = f'问题{question}_{STRATEGIES[kind]}搜索记录'
        FIGURE_NAMES[f'q{question}_summary_{kind}'] = f'问题{question}_{STRATEGIES[kind]}轨迹'

for scenario, label in SCENARIOS.items():
    prefix = 'q4_constant_reopt_' + scenario
    SEARCH_NAMES[prefix] = f'问题4：{label}恒功率重优化'
    FIGURE_NAMES[prefix + '_evaluations_search_history'] = f'问题4_{label}恒功率搜索记录'
    FIGURE_NAMES['Q4_dynamic_' + scenario] = f'问题4_{label}动态功率轨迹'

# ==================== calibrate ====================
def experiment(cfg, temp, data=None, label='q1'):
    d = data if data is not None else load_source()['experiments'][temp]
    cfg = cfg.changed(T0=temp + TM, ambient=temp + TM, horizon=float(d.t.iloc[-1]), charge_limit=1000000.0, current_limit=1000000.0, voltage_limit=0.0)
    protocol = Protocol('experiment', times=d.t.values, currents=d.j_A.values if cfg.basis == 'A' else d.j.values)
    return evaluate(cfg, protocol, label=label)

def waveform(result, d):
    valid = result['t'][-1] >= d.t.iloc[-1] - 1e-07 and result['metrics']['stop_reason'] == 'horizon'
    if not valid:
        return dict(complete=False, stop_reason=result['metrics']['stop_reason'], last_time=result['t'][-1])
    v = np.interp(d.t, result['t'], result['V'][:, 0])
    t = np.interp(d.t, result['t'], result['T'][:, 0])
    i = int(np.argmin(v))
    obs = d.loc[d.V == d.V.min(), 't']
    timeerr = max(obs.min() - d.t.iloc[i], d.t.iloc[i] - obs.max(), 0.0)
    depth = v[0] - v.min()
    recovery = v[-1] - v.min()
    metrics = dict(complete=True, V_RMSE=float(np.sqrt(np.mean((v - d.V) ** 2))), T_RMSE=float(np.sqrt(np.mean((t - d['T']) ** 2))), V_MAE=float(np.mean(abs(v - d.V))), T_MAE=float(np.mean(abs(t - d['T']))), V_min=float(v.min()), valley_time=float(d.t.iloc[i]), valley_time_error=float(timeerr), depth=float(depth), recovery=float(recovery), depth_error=float(abs(depth - (d.V.iloc[0] - d.V.min()))), recovery_error=float(abs(recovery - (d.V.iloc[-1] - d.V.min()))), dT=float(t[-1] - t[0]))
    metrics['prediction_pass'] = metrics['V_RMSE'] <= 0.03 and metrics['T_RMSE'] <= 0.3 and (timeerr <= 1) and (metrics['depth_error'] <= 0.03) and (metrics['recovery_error'] <= 0.03)
    return metrics

def calibrate_fit_job(basis='B', fixture='H1', train=(-20, -25), tier='P1', budget=12, mesh=(3, 2, 3, 3, 3)):
    data = load_source()['experiments']
    base = Config(basis=basis, fixture=fixture, mesh=mesh)
    xs = [np.log(0.3)] if tier == 'P1' else [np.log(0.3), np.log(1.0), np.log(1.0)]
    lower = [np.log(0.0001)] if tier == 'P1' else list(np.log([0.0001, 0.05, 0.05]))
    upper = [np.log(100.0)] if tier == 'P1' else list(np.log([100.0, 100.0, 100.0]))
    records = []
    wall = time.perf_counter()

    def unpack(x):
        return base.changed(j0=float(np.exp(x[0])), **{} if tier == 'P1' else dict(tau_b=float(np.exp(x[1])), tau_f=float(np.exp(x[2]))))

    def residual(x):
        cfg = unpack(x)
        residuals = []
        for temp in train:
            d = data[temp]
            r = experiment(cfg, temp, d, label='fit')
            complete = r['t'][-1] >= d.t.iloc[-1] - 1e-07 and r['metrics']['numerically_accepted']
            v = np.interp(d.t, r['t'], r['V'][:, 0])
            t = np.interp(d.t, r['t'], r['T'][:, 0])
            residuals.extend((v - d.V) / 0.03)
            residuals.extend((t - d['T']) / 0.3)
            if not complete:
                residuals[-2 * len(d):] = np.asarray(residuals[-2 * len(d):]) + 100 * (d.t.iloc[-1] - r['t'][-1] + 1)
            records.append(dict(temp=temp, j0=cfg.j0, tau_b=cfg.tau_b, tau_f=cfg.tau_f, complete=complete, run=r['metrics']['candidate_id'], wall_s=r['metrics']['wall_s']))
        return np.asarray(residuals)

    def jacobian(x):
        f0 = residual(x)
        columns = []
        for k in range(len(x)):
            xp = np.array(x, copy=True)
            step = 0.03 if x[k] + 0.03 < upper[k] else -0.03
            xp[k] += step
            columns.append((residual(xp) - f0) / step)
        return np.array(columns).T
    fit = least_squares(residual, xs, bounds=(lower, upper), jac=jacobian, max_nfev=budget, ftol=0.005, xtol=0.005, gtol=0.005, loss='soft_l1')
    cfg = unpack(fit.x)
    job = f'{basis}_{fixture}_{tier}_train' + '_'.join((str(v) for v in train))
    report = dict(job=job, train=train, tier=tier, config=cfg.to_dict(), cost=float(fit.cost), nfev=fit.nfev, optimizer_settings=dict(method='trf', loss='soft_l1', initial_log_parameters=xs, lower_log_bounds=lower, upper_log_bounds=upper, max_nfev=budget, log_jacobian_step=0.03, ftol=0.005, xtol=0.005, gtol=0.005, voltage_scale=0.03, temperature_scale=0.3), wall_s=time.perf_counter() - wall, optimizer_message=fit.message, bounds='exploratory; not measured confidence limits', boundary_hit=bool(np.any(abs(fit.x - np.asarray(lower)) < 0.05) | np.any(abs(fit.x - np.asarray(upper)) < 0.05)), results={})
    for temp, d in data.items():
        r = experiment(cfg, temp, d, label='q1')
        report['results'][str(temp)] = dict(role='fit' if temp in train else 'temperature_holdout', run=r['metrics']['candidate_id'], **waveform(r, d))
    J = fit.jac
    report['sensitivity_singular_values'] = np.linalg.svd(J, compute_uv=False).tolist()
    report['jacobian_note'] = 'optimizer finite-difference Jacobian; not a statistical confidence interval'
    write_json(OUTPUT / 'calibration' / f'{job}.json', report)
    pd.DataFrame(records).to_csv(OUTPUT / 'calibration' / f'{job}_evaluations.csv', index=False)
    return report

# ==================== completion ====================
def trajectory_diagnostics():
    rows = []
    mesh = []
    selected = _question_module('q2').completion_read('selected_model.json')
    src = load_source()['experiments']
    for job in [selected['joint_fit'], *selected['validation']]:
        report = _question_module('q2').completion_read('calibration/' + job + '.json')
        for temp, entry in report['results'].items():
            d = src[int(temp)]
            r = load_run(entry['run'])
            tt = d.t.values
            prediction = np.interp(tt, r['t'], r['T'][:, 0])
            observed = d['T'].values
            slopes = []
            for k, t in enumerate(tt):
                ids = (tt >= t - 1) & (tt <= t + 1) & (tt <= r['t'][-1])
                if ids.sum() < 3:
                    continue
                x = tt[ids] - tt[ids].mean()
                denom = x @ x
                so = x @ observed[ids] / denom
                sp = x @ prediction[ids] / denom
                slopes.append((t, so, sp))
            arr = np.array(slopes)
            pd.DataFrame(arr, columns=['t_s', 'observed_K_s', 'model_K_s']).to_csv(OUTPUT / f'slope_{job}_{temp}.csv', index=False)
            rows.append(dict(job=job, temp=int(temp), run=entry['run'], window_s=2, slope_RMSE_K_s=float(np.sqrt(np.mean((arr[:, 1] - arr[:, 2]) ** 2))), slope_max_error_K_s=float(np.max(abs(arr[:, 1] - arr[:, 2])))))
    for row in _question_module('q2').completion_read('mesh_convergence.json'):
        r = load_run(row['candidate_id'])
        m = Model(Config(**r['config']))
        peak = (-np.inf, None)
        for t, y in zip(r['t'], r['y']):
            s = m.phases(y)
            bulk = s['mi'] / 920
            idx = np.unravel_index(np.argmax(bulk), bulk.shape)
            value = float(bulk[idx])
            if value > peak[0]:
                peak = (value, (float(t), idx))
        value, (t, (cell, k)) = peak
        mesh.append(dict(run=row['candidate_id'], n_MEA=m.n, ice_bulk_max=value, t_peak_s=t, cell=int(cell + 1), layer=LAYER_NAMES[m.g.layer[k]], x_peak_m=float(m.g.x[k])))
    write_json(OUTPUT / 'temperature_slope_diagnostics.json', rows)
    pd.DataFrame(mesh).to_csv(OUTPUT / 'mesh_ice_peak.csv', index=False)
    return dict(slope_cases=len(rows), ice_peak_grids=len(mesh))

def failure_diagnostics():
    rows = []
    for path in OUTPUT.glob('q[23]_*_search.json'):
        groups = {}
        for candidate in _question_module('q2').completion_read(path.name)['candidates']:
            reason = candidate['stop_reason']
            if reason not in groups or candidate['last_valid_time'] > groups[reason]['last_valid_time']:
                groups[reason] = candidate
        for reason, candidate in groups.items():
            r = load_run(candidate['candidate_id'])
            m = Model(Config(**r['config']))
            y = r['y'][-1]
            s = m.phases(y)
            o = m.observations(s, r['j'][-1])
            ids = m.g.pem
            kappa = conductivity(s['local_T'][:, ids], s['lam'][:, ids])
            diffusivity = membrane_D(s['local_T'][:, ids], s['lam'][:, ids])
            for c in range(m.c):
                rows.append(dict(search=path.stem, selection='longest_valid_time_per_stop_reason', reason=reason, run=candidate['candidate_id'], cell=c + 1, t_s=float(r['t'][-1]), T_C=float(r['T'][-1, c]), V=float(o['V'][c]), activation_V=float(o['activation'][c]), ohmic_V=float(o['ohmic'][c]), concentration_V=float(o['concentration'][c]), lambda_mean=float(r['lambda_mean'][-1, c]), conductivity_min_S_m=float(kappa[c].min()), membrane_D_min_m2_s=float(diffusivity[c].min()), ice_bulk_max=float(o['ice_bulk'][c])))
            export_mechanisms(candidate['candidate_id'])
    pd.DataFrame(rows).to_csv(OUTPUT / 'failure_diagnostics.csv', index=False)
    return dict(cell_records=len(rows), scope='Membrane constitutive-domain exits do not prove a real heater strategy impossible')

# ==================== preflight ====================
def preflight():
    src = audit()
    OUTPUT.mkdir(parents=True, exist_ok=True)
    heats = []
    water = []
    volts = []
    stack = []
    for temp, d in src['experiments'].items():
        for basis in ('A', 'B'):
            j = d.j_A.values if basis == 'A' else d.j.values
            q = np.trapezoid(j, d.t)
            Q = 25 * np.trapezoid(j * (1.48 - d.V), d.t)
            heats.append(dict(T0_C=temp, basis=basis, Q_J=Q, dT_K=d['T'].iloc[-1] - d['T'].iloc[0], C_app=Q / (d['T'].iloc[-1] - d['T'].iloc[0])))
            water.append(dict(T0_C=temp, basis=basis, charge_C_cm2=q, water_mg_cm2=q * MW / (2 * F) * 1000000.0, pore_cCL_m3_m2=THICKNESS[3] * POROSITY[3], liquid_capacity_C_cm2=990 * THICKNESS[3] * POROSITY[3] * 2 * F / MW / 10000.0))
            for index in (0, len(d) - 1):
                for j0 in (0.01, 1.0, 1e+20):
                    m = Model(Config(T0=float(d['T'].iloc[index]) + TM, j0=j0, basis=basis))
                    o = m.voltage(m.phases(m.initial()), j[index])
                    volts.append(dict(T0_C=temp, basis=basis, t_s=float(d.t.iloc[index]), j0=j0, V=float(o['V'][0]), observed_V=float(d.V.iloc[index])))
    for shared in (False, True):
        for thermal in ('effective', 'phase'):
            m = Model(Config(cells=5, fixture='H2', shared_bp=shared, thermal=thermal))
            k = m.thermal_k(m.phases(m.initial()))
            g = m.g
            C = np.dot(g.C, g.tdx)
            L = sum(g.tdx)
            ke = L / sum(g.tdx / k)
            stack.append(dict(BP_count=6 if shared else 10, thermal=thermal, C_J_m2K=C, L_m=L, k_eff=ke, Bi=40 * L / 2 / ke, tau_min=C / 80 / 60))
    for name, rows in [('heat', heats), ('water', water), ('voltage', volts), ('stack', stack)]:
        pd.DataFrame(rows).to_csv(OUTPUT / f'preflight_{name}.csv', index=False)
    cases = [Config(basis=b, fixture=h).to_dict() for b, h in [('B', 'H1'), ('A', 'H2'), ('A', 'H1'), ('B', 'H2')]]
    write_json(io_ROOT / 'configs' / 'case_registry.json', dict(cases=cases, source_sha256=src['sha256'], fixed_choices=dict(temperature='MEA spatial volume mean', q34_charge='no cap; 20 C/cm2 sensitivity', water='W1 primary; W0 ablation', outlet='closed primary; pressure-release drainage envelope', precool='fixed water inventories; reinitialize dry free-water at startup', gas_sensible='neglected; dry materials plus water enthalpy'), limitations=['low-temperature capillary extrapolation', 'experimental current basis unresolved', 'no independent ice data']))
    return src

# ==================== verification ====================
def water_job(args):
    cfg, temp, label = args
    if temp in (-20, -25):
        r = experiment(cfg, temp, label=label)
    else:
        case = cfg.changed(cells=5, fixture='H2', T0=temp + TM, ambient=temp + TM, horizon=100, charge_limit=20 if temp == -10 else 1000000.0)
        p = Protocol('constant', (0.2,)) if temp == -10 else Protocol('fixed', powers=(1.0,) * 5, heat_off=80)
        r = evaluate(case, p, label=label, stop_success=temp == -10)
    return dict(temperature=temp, water=cfg.water, outlet=cfg.outlet, **r['metrics'])

def verify_water(cfg, workers=4):
    jobs = [(cfg.changed(water=w, outlet=o), t, 'water_route') for w, o in [('W0', 'closed'), ('W1', 'closed'), ('W1', 'drain')] for t in (-20, -25, -10, -30)]
    with ProcessPoolExecutor(max_workers=workers) as pool:
        rows = list(pool.map(water_job, jobs))
    write_json(OUTPUT / 'water_route_ablation.json', rows)
    pd.DataFrame([{k: v for k, v in r.items() if not isinstance(v, (dict, list))} for r in rows]).to_csv(OUTPUT / 'water_route_ablation.csv', index=False)
    (OUTPUT / 'water_model_release.md').write_text('# 水通路放行记录\n\nW1 继续保留。W0/W1 的实验、−10℃自启动及−30℃加热结果见 water_route_ablation.csv。\n\n尚未证明全参数集合的 ts/E/Vmin/最大冰阈值全部等效，故不批准把 W0 简化为最终模型。低温毛细本构与流道排液边界作为条件假设报告；数值守恒不等于实验有效性。\n', encoding='utf-8')
    return rows

def convergence(cfg, temp=-20):
    rows = []
    runs = []
    for mesh, step, rtol in [((3, 2, 3, 3, 3), 0.5, 1e-05), ((12, 6, 8, 10, 12), 0.2, 2e-06), ((24, 12, 16, 20, 24), 0.1, 1e-06), ((48, 24, 32, 40, 48), 0.05, 5e-07)]:
        r = experiment(cfg.changed(mesh=mesh, max_step=step, rtol=rtol, board_n=4, end_n=8), temp, label='convergence')
        runs.append(r)
        rows.append(dict(n_MEA=sum(mesh), max_step=step, rtol=rtol, **r['metrics']))
    d = load_source()['experiments'][temp]
    for i in range(len(runs) - 1):
        a, b = runs[i:i + 2]
        end = min(a['t'][-1], b['t'][-1])
        tt = d.t[d.t <= end].values
        dv = np.max(abs(np.interp(tt, a['t'], a['V'][:, 0]) - np.interp(tt, b['t'], b['V'][:, 0])))
        dt = np.max(abs(np.interp(tt, a['t'], a['T'][:, 0]) - np.interp(tt, b['t'], b['T'][:, 0])))
        rows[i]['next_grid_V_diff'] = float(dv)
        rows[i]['next_grid_T_diff'] = float(dt)
        rows[i]['output_converged'] = bool(dv < 0.001 and dt < 0.02 and (end >= 36.6 - 1e-07))
    if not rows[-2]['output_converged']:
        last = Config(**runs[-1]['config'])
        extra = experiment(last.changed(mesh=tuple((2 * n for n in last.mesh)), max_step=last.max_step / 2, rtol=last.rtol / 2), temp, label='convergence_extra')
        a, b = (runs[-1], extra)
        end = min(a['t'][-1], b['t'][-1])
        tt = d.t[d.t <= end].values
        dv = float(np.max(abs(np.interp(tt, a['t'], a['V'][:, 0]) - np.interp(tt, b['t'], b['V'][:, 0]))))
        dt = float(np.max(abs(np.interp(tt, a['t'], a['T'][:, 0]) - np.interp(tt, b['t'], b['T'][:, 0]))))
        rows[-1].update(next_grid_V_diff=dv, next_grid_T_diff=dt, output_converged=bool(dv < 0.001 and dt < 0.02 and (end >= 36.6 - 1e-07)))
        rows.append(dict(n_MEA=sum(extra['config']['mesh']), max_step=extra['config']['max_step'], rtol=extra['config']['rtol'], **extra['metrics']))
    write_json(OUTPUT / 'mesh_convergence.json', rows)
    pd.DataFrame([{k: v for k, v in r.items() if not isinstance(v, (dict, list))} for r in rows]).to_csv(OUTPUT / 'mesh_convergence.csv', index=False)
    return rows

def export_mechanisms(run_name):
    r = load_run(run_name)
    cfg = Config(**r['config'])
    m = Model(cfg)
    rows = []
    field = []
    for n, (t, y, j, u) in enumerate(zip(r['t'], r['y'], r['j'], r['u'])):
        s = m.phases(y)
        d = m.rhs(t, y, j, u, True)
        o = m.observations(s, j)
        for c in range(m.c):
            for layer, name in enumerate(LAYER_NAMES):
                ix = np.flatnonzero(m.g.layer == layer)
                dx = m.g.dx[ix]
                a, b = (ix[0], ix[-1] + 1)
                rows.append(dict(t_s=t, cell=c + 1, layer=name, vapor_kg=float(np.dot(s['mv'][c, ix], dx) * AREA), liquid_kg=float(np.dot(s['ml'][c, ix], dx) * AREA), ice_kg=float(np.dot(s['mi'][c, ix], dx) * AREA), bound_kg=float(np.dot(s['mb'][c, ix], dx) * AREA), production_kg_s=MW * j * 10000.0 * AREA / (2 * F) if layer == 3 else 0.0, **{name + '_net_out_kg_s': float((f[c, b] - f[c, a]) * AREA) for name, f in d['fluxes'].items() if name != 'gas'}, reaction_W=float(d['reaction_W_m2'][c] * AREA), activation_V=float(o['activation'][c]), ohmic_V=float(o['ohmic'][c]), concentration_V=float(o['concentration'][c])))
        if n in (0, len(r['t']) // 3, len(r['t']) // 2, len(r['t']) - 1):
            for c in range(m.c):
                for k, x in enumerate(m.g.x):
                    field.append(dict(t_s=t, cell=c + 1, x_m=x, layer=LAYER_NAMES[m.g.layer[k]], T_K=s['local_T'][c, k], mv=s['mv'][c, k], ml=s['ml'][c, k], mi=s['mi'][c, k], mb=s['mb'][c, k], lambda_=s['lam'][c, k], epsilon_g=s['eg'][c, k]))
    pd.DataFrame(rows).to_csv(_project_storage() / 'runs' / run_name / 'water_budget_by_layer.csv', index=False)
    pd.DataFrame(field).to_csv(_project_storage() / 'runs' / run_name / 'spatial_snapshots.csv', index=False)

# ==================== uncertainty ====================
def sensitivity_job(args):
    name, cfg = args
    data = load_source()['experiments']
    rows = []
    for t, d in data.items():
        r = experiment(cfg, t, label='sensitivity')
        rows.append(dict(branch=name, T0_C=t, run=r['metrics']['candidate_id'], **waveform(r, d)))
    return rows

def uncertainty(cfg, workers=4):
    alternatives = [('baseline', cfg), ('tau_b_0.1x', cfg.changed(tau_b=cfg.tau_b * 0.1)), ('tau_b_10x', cfg.changed(tau_b=cfg.tau_b * 10)), ('tau_f_0.1x', cfg.changed(tau_f=cfg.tau_f * 0.1)), ('tau_f_10x', cfg.changed(tau_f=cfg.tau_f * 10)), ('mu_2x', cfg.changed(mu_factor=2)), ('mu_5x', cfg.changed(mu_factor=5)), ('ice_connectivity_1', cfg.changed(ice_connectivity=1)), ('ice_connectivity_5', cfg.changed(ice_connectivity=5)), ('effective_k', cfg.changed(thermal='effective')), ('unit_temperature', cfg.changed(observation='unit')), ('unified_concentration', cfg.changed(concentration='unified')), ('half_CL_path', cfg.changed(half_cl=True)), ('h_20', cfg.changed(h=20)), ('h_60', cfg.changed(h=60))]
    with ProcessPoolExecutor(max_workers=workers) as pool:
        nested = list(pool.map(sensitivity_job, alternatives))
    rows = [row for group in nested for row in group]
    write_json(OUTPUT / 'sensitivity.json', rows)
    pd.DataFrame(rows).to_csv(OUTPUT / 'sensitivity.csv', index=False)
    return rows

# ==================== identifiability ====================
def response_job(cfg):
    data = load_source()['experiments']
    outputs = []
    ids = []
    for temp, d in data.items():
        r = experiment(cfg, temp, d, label='identifiability')
        ids.append(r['metrics']['candidate_id'])
        if r['t'][-1] < d.t.iloc[-1] - 1e-07 or not r['metrics']['numerically_accepted']:
            return (None, ids)
        outputs.extend(np.interp(d.t, r['t'], r['V'][:, 0]) / 0.03)
        outputs.extend(np.interp(d.t, r['t'], r['T'][:, 0]) / 0.3)
    return (np.array(outputs), ids)

def identifiability(cfg, workers=4):
    parameters = ['j0', 'tau_b', 'tau_f']
    steps = [0.03, 0.015]
    jobs = [cfg] + [cfg.changed(**{key: getattr(cfg, key) * np.exp(h)}) for h in steps for key in parameters]
    with ProcessPoolExecutor(max_workers=workers) as pool:
        results = list(pool.map(response_job, jobs))
    if any((r[0] is None for r in results)):
        report = dict(status='unresolved', reason='at least one perturbation lacks a complete physical trajectory')
    else:
        base = results[0][0]
        matrices = []
        for k, h in enumerate(steps):
            matrices.append(np.array([(results[1 + k * 3 + i][0] - base) / h for i in range(3)]).T)
        change = np.linalg.norm(matrices[0] - matrices[1], axis=0) / np.maximum(np.linalg.norm(matrices[1], axis=0), 1e-12)
        sv = np.linalg.svd(matrices[1], compute_uv=False)
        norms = np.linalg.norm(matrices[1], axis=0)
        corr = matrices[1].T @ matrices[1] / np.maximum(norms[:, None] * norms[None, :], 1e-20)
        report = dict(status='computed', parameters=parameters, log_steps=steps, relative_derivative_change=change, singular_values=sv, condition_number=float(sv[0] / sv[-1]) if sv[-1] > 0 else None, sensitivity_column_cosines=corr, statistical_interpretation='local sensitivity only; model misspecification and serial dependence preclude confidence claims')
        pd.DataFrame(matrices[1], columns=parameters).to_csv(OUTPUT / 'parameter_sensitivity_matrix.csv', index=False)
    report['runs'] = [r[1] for r in results]
    write_json(OUTPUT / 'identifiability.json', report)
    return report

# ==================== tables ====================
def make_tables():
    src = load_source()
    tables = OUTPUT / 'tables'
    tables.mkdir(parents=True, exist_ok=True)
    selected_path = OUTPUT / 'selected_model.json'
    if selected_path.exists():
        selected = json.loads(selected_path.read_text(encoding='utf-8'))
        joint = json.loads((OUTPUT / 'calibration' / (selected['joint_fit'] + '.json')).read_text(encoding='utf-8'))
        for no, temp in enumerate((-20, -25)):
            result = joint['results'][str(temp)]
            r = load_run(result['run'])
            rows = []
            for row in src['tables'][no][1:]:
                if not row or not row[0].strip():
                    continue
                try:
                    t = float(row[0])
                    v = float(row[1])
                    tempobs = float(row[4])
                except (ValueError, IndexError):
                    continue
                valid = t <= r['t'][-1] + 1e-08
                vs = float(np.interp(t, r['t'], r['V'][:, 0])) if valid else np.nan
                ts = float(np.interp(t, r['t'], r['T'][:, 0])) if valid else np.nan
                ice = float(np.interp(t, r['t'], r['ice'][:, 0])) if valid else np.nan
                rows.append(dict(time_s=t, experiment_V=v, model_V=vs, V_relative_error_pct=100 * abs(vs - v) / abs(v), experiment_T_C=tempobs, model_T_C=ts, T_relative_error_pct=100 * abs(ts - tempobs) / abs(tempobs), max_ice_bulk=ice, status='joint_fit_not_independent_validation' if valid else 'after_physical_failure', run=result['run']))
            pd.DataFrame(rows).to_csv(tables / f'table{no + 1}_Q1_{temp}.csv', index=False, encoding='utf-8-sig')
    for filename, outname in [('q2_summary.json', 'table3_Q2.csv'), ('q3_summary.json', 'table4_Q3.csv')]:
        p = OUTPUT / filename
        if filename == 'q3_summary.json' and (OUTPUT / 'q3_extended_summary.json').exists():
            p = OUTPUT / 'q3_extended_summary.json'
        if p.exists():
            data = json.loads(p.read_text(encoding='utf-8'))
            rows = []
            for kind, w in data.items():
                if w is None:
                    prefix = 'q2' if filename.startswith('q2') else 'q3'
                    search = json.loads((OUTPUT / f'{prefix}_{kind}_search.json').read_text(encoding='utf-8'))['candidates']
                    from collections import Counter
                    reasons = Counter((r['stop_reason'] for r in search))
                    rows.append(dict(strategy=kind, status='no_feasible_candidate_found', startup_time_s=np.nan, E_aux_J=np.nan, evaluated=len(search), failure_reasons=json.dumps(reasons), last_valid_min_s=min((r['last_valid_time'] for r in search)), last_valid_max_s=max((r['last_valid_time'] for r in search))))
                    continue
                fine = w.get('second_fine_replay', w.get('fine_replay', {}))
                reported = fine or w
                rows.append(dict(strategy=kind, status=reported['feasibility'], startup_time_s=reported['success_time'], heat_off_s=reported['th_startup'], E_aux_J=reported['E_aux_total_J'], charge_C_cm2=reported['charge_C_cm2'], minimum_V=reported['V_min'], maximum_ice_bulk=reported['ice_max'], maximum_cell_delta_T=reported['delta_T_max'], search_startup_s=w['success_time'], search_run=w['candidate_id'], control_expression=json.dumps(w['protocol']), run=reported['candidate_id'], fine_status=fine.get('feasibility', 'not_run'), fine_startup_s=fine.get('success_time'), fine_energy_J=fine.get('E_aux_total_J'), qualification_status=w.get('qualification', {}).get('stop_reason', 'not_applicable'), qualification_pass=w.get('qualification', {}).get('qualification_pass')))
            pd.DataFrame(rows).to_csv(tables / outname, index=False, encoding='utf-8-sig')
    p = OUTPUT / 'q4_extended_comparison.json'
    if not p.exists():
        p = OUTPUT / 'q4_comparison.json'
    if p.exists():
        rows = json.loads(p.read_text(encoding='utf-8'))
        fine_path = OUTPUT / 'q4_fine_replay.json'
        checked = dict(zip(('equilibrium', '20min', '40min'), json.loads(fine_path.read_text(encoding='utf-8')))) if fine_path.exists() else {}
        for row in rows:
            if row['method'].startswith('dynamic'):
                row['power_control_expression'] = 'clip(ff+KT*max(Ttarget-Tf,0)+KV*max(Vwarn-Vf,0)-Kh*max(Tf-Toff,0)+Kr*max(rmin-dTf/dt,0)+Kd*max(-dR/dt-dwarn,0),0,1); R=Vf-Vref; invalid-reference fallback R=Vf; derivative terms disabled before valid history; latched 0 after confirmed shutdown; gains in rule'
            if row['method'].startswith('dynamic') and row['scenario'] in checked:
                replay = checked[row['scenario']]
                row.update(fine_status=replay['feasibility'], fine_startup_s=replay['success_time'], fine_energy_J=replay['E_aux_total_J'], fine_run=replay['candidate_id'])
            if 'shutdown_extrema' in row:
                for key, value in row['shutdown_extrema'].items():
                    row[key + '_to_official_success'] = row.get(key)
                    row[key] = value
                row['extrema_window'] = 'start_to_control_exit'
        pd.DataFrame([{k: json.dumps(v) if isinstance(v, (dict, list)) else v for k, v in row.items()} for row in rows]).to_csv(tables / 'table4_Q4.csv', index=False, encoding='utf-8-sig')
    ledger = []
    for p in (_project_storage() / 'evaluation_ledger').glob('*.json'):
        r = json.loads(p.read_text(encoding='utf-8'))
        ledger.append({k: v for k, v in r.items() if not isinstance(v, (dict, list))})
    pd.DataFrame(ledger).to_csv(OUTPUT / 'evaluation_ledger.csv', index=False)

def parameter_registry():
    raw = load_source()['parameters']
    rows = []
    for row in raw:
        value = row['raw_value']
        si = value
        factor = 1.0
        if isinstance(value, str):
            try:
                si = float(value.strip())
            except ValueError:
                pass
        if row['source_cell'].endswith('!C2'):
            factor = 0.0001
        if row['source_cell'].endswith('!C27'):
            factor = 0.001
        if isinstance(si, (int, float)):
            si *= factor
        unit = 'm^2' if row['source_cell'].endswith('!C2') else 'kg/mol' if row['source_cell'].endswith('!C27') else row['unit']
        rows.append(dict(parameter=row['name'], value=si, raw_value=value, unit=unit, raw_unit=row['unit'], source=row['source_cell'], status='given' if row['source_cell'] not in ('参数清单!C46', '参数清单!C47', '参数清单!C48', '参数清单!C49', '参数清单!C50') else 'dimensionless_not_used_as_rate', conversion_factor=factor))
        if isinstance(value, str) and ';' in value and (';' in str(row['unit'])):
            values = [float(x.strip()) for x in value.split(';')]
            units = [x.strip() for x in row['unit'].split(';')]
            for component, v, u in zip(('density', 'specific_heat', 'thermal_conductivity', 'electric_conductivity'), values, units):
                rows.append(dict(parameter=row['name'] + '.' + component, value=v, unit=u, source=row['source_cell'], status='SI_component_parsed_from_source', conversion_factor=1.0))
    for key, value in Config().to_dict().items():
        rows.append(dict(parameter='config.' + key, value=str(value), unit='see Config and README', source='implementation_registry', status='registered_choice'))
    for row in rows:
        row['bounds'] = 'fixed/not fitted'
    selected_path = OUTPUT / 'selected_model.json'
    if selected_path.exists():
        selected = json.loads(selected_path.read_text(encoding='utf-8'))
        units = {'j0': 'A/m^2', 'tau_b': 's', 'tau_f': 's at -20 C', 'tau_m': 's at 5 C', 'h': 'W/(m^2 K)', 'Rc': 'ohm m^2', 'T0': 'K', 'ambient': 'K', 'lambda0': '1', 'mu_factor': '1', 'ice_connectivity': '1'}
        ranges = {'j0': '[0.0001, 100] exploratory', 'tau_b': '[0.05, 100] exploratory', 'tau_f': '[0.05, 100] exploratory'}
        for key, value in selected['config'].items():
            rows.append(dict(parameter='frozen.' + key, value=str(value), unit=units.get(key, 'see Config and README'), source='selected_model.json / ' + selected['joint_fit'], status='frozen_conditional_model', bounds=ranges.get(key, 'fixed')))
    (io_ROOT / 'configs').mkdir(exist_ok=True)
    pd.DataFrame(rows).to_csv(io_ROOT / 'configs' / 'parameter_registry.csv', index=False, encoding='utf-8-sig')
    pd.DataFrame([dict(parameter='j0', equation='Arrhenius/asinh activation', observable='voltage', confounder='ice activity, temperature', status='P1'), dict(parameter='tau_b', equation='bound-water relaxation', observable='voltage transient', confounder='freeze time', status='P3 exploratory 0.05..100 s'), dict(parameter='tau_f', equation='freeze at -20 C', observable='voltage/temperature transient', confounder='drainage, adsorption', status='P3 exploratory 0.05..100 s')]).to_csv(OUTPUT / 'parameter_observation_map.csv', index=False)

# ==================== figures ====================
matplotlib.use('Agg')

figures_COLORS = ['#4477AA', '#EE6677', '#228833', '#CCBB44', '#AA3377']

configure_chinese_plotting()

def save(fig, name):
    directory = OUTPUT / 'figures'
    directory.mkdir(exist_ok=True)
    fig.savefig(directory / f'{name}.svg', bbox_inches='tight')
    fig.savefig(directory / f'{name}.pdf', bbox_inches='tight')
    fig.savefig(directory / f'{name}.png', dpi=300, bbox_inches='tight')
    plt.close(fig)

def figures():
    sources = []
    p = OUTPUT / 'selected_model.json'
    if p.exists():
        selected = json.loads(p.read_text(encoding='utf-8'))
        fit = json.loads((OUTPUT / 'calibration' / (selected['joint_fit'] + '.json')).read_text(encoding='utf-8'))
        src = load_source()['experiments']
        fig, axes = plt.subplots(2, 2, figsize=(7.2, 4.8), layout='constrained')
        for k, temp in enumerate((-20, -25)):
            r = load_run(fit['results'][str(temp)]['run'])
            d = src[temp]
            axes[0, k].plot(d.t, d.V, color='#777777', lw=1, label='实验数据（184点）')
            axes[0, k].plot(r['t'], r['V'][:, 0], color=figures_COLORS[0], label='联合校准')
            axes[1, k].plot(d.t, d['T'], color='#777777', lw=1, label='实验数据')
            axes[1, k].plot(r['t'], r['T'][:, 0], color=figures_COLORS[0], label='联合校准')
            axes[0, k].set(title=f'{temp} ℃', ylabel='单片电压（V）')
            axes[1, k].set(xlabel='时间（s）', ylabel='平均温度（℃）')
            sources.append(fit['results'][str(temp)]['run'])
        axes[0, 0].legend(fontsize=7)
        fig.suptitle('模型校准与实验对比')
        save(fig, 'Q1_calibration')
        fig, axes = plt.subplots(1, 2, figsize=(7.2, 2.6), layout='constrained')
        for k, temp in enumerate((-20, -25)):
            r = load_run(fit['results'][str(temp)]['run'])
            axes[0].plot(r['t'], r['ice'][:, 0], color=figures_COLORS[k], label=f'{temp} ℃')
            axes[1].plot(r['t'], r['lambda_mean'][:, 0], color=figures_COLORS[k], label=f'{temp} ℃')
        axes[0].set(xlabel='时间（s）', ylabel='最大冰体积分数（模型推断）')
        axes[1].set(xlabel='时间（s）', ylabel='平均膜含水量 λ')
        axes[0].legend()
        save(fig, 'Q1_inferred_states')
        fig, axes = plt.subplots(2, 2, figsize=(7.2, 4.8), layout='constrained')
        for k, job in enumerate(selected['validation']):
            v = json.loads((OUTPUT / 'calibration' / (job + '.json')).read_text(encoding='utf-8'))
            temp = next((t for t in (-20, -25) if t not in v['train']))
            r = load_run(v['results'][str(temp)]['run'])
            d = src[temp]
            for ax, field, obs, label in [(axes[0, k], 'V', d.V, '电压（V）'), (axes[1, k], 'T', d['T'], '温度（℃）')]:
                ax.plot(d.t, obs, color='#777777', label='实验数据')
                ax.plot(r['t'], r[field][:, 0], color=figures_COLORS[1], label='留出工况预测')
                ax.set(xlabel='时间（s）', ylabel=label)
            axes[0, k].set_title(f"训练：{v['train'][0]} ℃ → 检验：{temp} ℃")
        axes[0, 0].legend(fontsize=7)
        save(fig, 'Q1_temperature_holdout')
    p = OUTPUT / 'precool_dense.csv'
    if p.exists():
        d = pd.read_csv(p)
        fig, axes = plt.subplots(1, 2, figsize=(7.2, 2.8), layout='constrained')
        for i in range(5):
            axes[0].plot(d.minutes, d[f'T{i + 1}_C'], color=figures_COLORS[i], label=f'第{i + 1}片')
        axes[1].plot(d.minutes, d.delta_T_K, color=figures_COLORS[0])
        axes[0].set(xlabel='预冷时间（min）', ylabel='单片平均温度（℃）')
        axes[1].set(xlabel='预冷时间（min）', ylabel='片间温差（K）')
        axes[0].legend(fontsize=6)
        save(fig, 'Q4_precooling')
        data = np.load(OUTPUT / 'precool_fields.npz')
        fig, ax = plt.subplots(figsize=(7.2, 2.6), layout='constrained')
        for i, minute in enumerate((20, 40, 100)):
            k = np.argmin(abs(data['minutes'] - minute))
            ax.plot(data['x'] * 1000, data['T'][k] - 273.15, color=figures_COLORS[i], label=f'预冷{minute} min')
        ax.set(xlabel='堆叠方向位置（mm）', ylabel='温度（℃）')
        ax.legend()
        save(fig, 'Q4_precooling_spatial')
    for filename in ('q2_summary', 'q3_summary'):
        p = OUTPUT / (filename + '.json')
        if filename == 'q3_summary' and (OUTPUT / 'q3_extended_summary.json').exists():
            p = OUTPUT / 'q3_extended_summary.json'
        if not p.exists():
            continue
        rows = json.loads(p.read_text(encoding='utf-8'))
        for kind, w in rows.items():
            representative = False
            if not w:
                prefix = filename.split('_')[0]
                search = OUTPUT / f'{prefix}_{kind}_search.json'
                if not search.exists():
                    continue
                candidates = json.loads(search.read_text(encoding='utf-8'))['candidates']
                w = max(candidates, key=lambda x: x['last_valid_time'])
                representative = True
            chosen = w.get('second_fine_replay', w.get('fine_replay', w))
            r = load_run(chosen['candidate_id'])
            sources.append(chosen['candidate_id'])
            fig, axes = plt.subplots(2, 2, figsize=(7.2, 4.8), layout='constrained')
            for i in range(5):
                for ax, field, label in [(axes[0, 0], 'T', '单片温度（℃）'), (axes[0, 1], 'V', '单片电压（V）'), (axes[1, 0], 'ice', '最大冰体积分数'), (axes[1, 1], 'u', '加热功率密度（W/cm²）')]:
                    ax.plot(r['t'], r[field][:, i], color=figures_COLORS[i], label=f'第{i + 1}片')
                    ax.set(xlabel='时间（s）', ylabel=label)
                    if field == 'u':
                        ax.set_ylim(0, 1.05)
            axes[0, 1].axhline(0.3, color='#555555', ls='--', lw=0.8)
            axes[1, 0].set_ylim(0, max(0.001, float(np.max(r['ice'])) * 1.1))
            if filename == 'q2_summary':
                axes[1, 1].clear()
                axes[1, 1].plot(r['t'], r['j'], color=figures_COLORS[0])
                axes[1, 1].set(xlabel='时间（s）', ylabel='电流密度（A/cm²）', ylim=(0, 0.5))
            axes[0, 0].legend(fontsize=6)
            suffix = f"——代表性{STATUSES[w['feasibility']]}轨迹" if representative else '（条件性模型结果）'
            fig.suptitle(f'问题{filename[1]}：{STRATEGIES[kind]}' + suffix)
            save(fig, filename + '_' + kind)
    for path in OUTPUT.glob('*_evaluations.csv'):
        d = pd.read_csv(path)
        if 'best_so_far' not in d:
            continue
        fig, axes = plt.subplots(1, 2, figsize=(7.2, 2.8), layout='constrained')
        valid = np.isfinite(d.best_so_far)
        if valid.any():
            axes[0].step(d.evaluation[valid], d.best_so_far[valid], where='post', color=figures_COLORS[0])
        else:
            axes[0].text(0.5, 0.5, '未找到可行候选', transform=axes[0].transAxes, ha='center')
            axes[0].set_axis_off()
        objective = '当前最优可行启动时间（s）' if path.name.startswith('q2') else '当前最优可行加热能耗（J）'
        axes[0].set(xlabel='候选评价次数', ylabel=objective, title=SEARCH_NAMES[path.stem.removesuffix('_evaluations')])
        if path.name.startswith(('q3_', 'q4_')):
            axes[0].set_title(SEARCH_NAMES[path.stem.removesuffix('_evaluations')] + '\n原300 s时域搜索')
        counts = d.stop_reason.value_counts()
        axes[1].barh([STOP_REASONS[x] for x in counts.index], counts.values, color=figures_COLORS[0])
        axes[1].set(xlabel='候选数量', title='计算终止原因')
        save(fig, path.stem + '_search_history')
    p = OUTPUT / 'q4_extended_comparison.json'
    if not p.exists():
        p = OUTPUT / 'q4_comparison.json'
    if p.exists():
        rows = json.loads(p.read_text(encoding='utf-8'))
        for row in rows:
            if not row['method'].startswith('dynamic'):
                continue
            r = load_run(row['candidate_id'])
            fig, axes = plt.subplots(2, 2, figsize=(7.2, 4.8), layout='constrained')
            log = json.loads((_project_storage() / 'runs' / row['candidate_id'] / 'controller_log.json').read_text(encoding='utf-8'))
            log = [entry for entry in log if entry['t'] <= r['t'][-1] + 1e-08]
            for i in range(5):
                for ax, field, label in [(axes[0, 0], 'T', '平均温度（℃）'), (axes[0, 1], 'V', '电压（V）'), (axes[1, 0], 'ice', '最大冰体积分数'), (axes[1, 1], 'u', '加热功率密度（W/cm²）')]:
                    if field == 'u':
                        ax.step([entry['t'] for entry in log], [entry['u'][i] for entry in log], where='post', color=figures_COLORS[i], label=f'第{i + 1}片')
                    else:
                        ax.plot(r['t'], r[field][:, i], color=figures_COLORS[i], label=f'第{i + 1}片')
                    ax.set(xlabel='时间（s）', ylabel=label)
                    if field == 'u':
                        ax.set_ylim(0, 1.05)
            axes[0, 1].axhline(0.3, color='#555555', ls='--', lw=0.8)
            axes[0, 0].legend(fontsize=6)
            axes[1, 0].set_ylim(0, max(0.001, float(np.max(r['ice'])) * 1.1))
            fig.suptitle(f"问题4：{SCENARIOS[row['scenario']]}，{STRATEGIES[row['method']]}")
            save(fig, 'Q4_dynamic_' + row['scenario'])
    p = OUTPUT / 'q4_extended_scan.json'
    if not p.exists():
        p = OUTPUT / 'q4_precool_refined.json'
    if not p.exists():
        p = OUTPUT / 'q4_precool_constant_scan.json'
    if p.exists():
        d = pd.DataFrame(json.loads(p.read_text(encoding='utf-8'))).sort_values('precool_minutes')
        fig, axes = plt.subplots(1, 2, figsize=(7.2, 2.8), layout='constrained')
        axes[0].plot(d.precool_minutes, d.delta_T_max, color=figures_COLORS[0], marker='o', ms=3)
        axes[0].set(xlabel='预冷时间（min）', ylabel='启动过程最大片间温差（K）')
        for status, part in d.groupby('feasibility'):
            axes[1].scatter(part.precool_minutes, part.last_valid_time, label=STATUSES[status], s=15)
        axes[1].set(xlabel='预冷时间（min）', ylabel='轨迹最后有效时刻（s）')
        axes[1].legend(fontsize=6)
        save(fig, 'Q4_precool_startup_scan')
    p = OUTPUT / 'mesh_convergence.csv'
    if p.exists():
        d = pd.read_csv(p)
        fig, axes = plt.subplots(1, 2, figsize=(7.2, 2.8), layout='constrained')
        axes[0].loglog(d.n_MEA, d.next_grid_V_diff * 1000, 'o-', color=figures_COLORS[0])
        axes[0].axhline(1, color='#777777', ls='--')
        axes[0].set(xlabel='膜电极网格数', ylabel='相邻两级网格电压差（mV）')
        axes[1].loglog(d.n_MEA, d.next_grid_T_diff, 'o-', color=figures_COLORS[1])
        axes[1].axhline(0.02, color='#777777', ls='--')
        axes[1].set(xlabel='膜电极网格数', ylabel='相邻两级网格温度差（K）')
        ticks = d.loc[d.next_grid_V_diff.notna(), 'n_MEA'].astype(int).tolist()
        for ax in axes:
            ax.set_xticks(ticks, labels=[str(v) for v in ticks])
            ax.xaxis.set_minor_locator(matplotlib.ticker.NullLocator())
        save(fig, 'grid_convergence')
    p = OUTPUT / 'q3_coheat_evaluations.csv'
    if p.exists():
        d = pd.read_csv(p)
        d = d[d.feasibility == 'feasible'].sort_values('success_time')
        if (OUTPUT / 'q3_extended_summary.json').exists():
            zero = json.loads((OUTPUT / 'q3_extended_summary.json').read_text(encoding='utf-8'))['coheat']
            d = pd.concat([d, pd.DataFrame([zero])], ignore_index=True).sort_values('success_time')
        fig, ax = plt.subplots(figsize=(3.6, 2.8), layout='constrained')
        if len(d):
            ax.scatter(d.success_time, d.E_aux_total_J, s=12, color=figures_COLORS[0], alpha=0.65, label='搜索网格上的可行候选')
            minimum = np.minimum.accumulate(d.E_aux_total_J.values)
            take = np.r_[True, np.diff(minimum) < 0]
            ax.plot(d.success_time.values[take], minimum[take], color=figures_COLORS[1], marker='o', ms=3, label='采样非支配解集')
            ax.legend(fontsize=6)
        ax.set(xlabel='启动时间（s）', ylabel='辅助加热能耗（J）', title='条件性模型的能耗与时间权衡')
        save(fig, 'Q3_energy_time_tradeoff')
    p = OUTPUT / 'sensitivity.csv'
    if p.exists():
        d = pd.read_csv(p)
        fig, axes = plt.subplots(1, 2, figsize=(9, 4.8), layout='constrained')
        for ax, field, label in [(axes[0], 'V_RMSE', '电压均方根误差（V）'), (axes[1], 'T_RMSE', '温度均方根误差（K）')]:
            for k, (temp, part) in enumerate(d.groupby('T0_C')):
                ax.scatter(part[field], [SENSITIVITY[x] for x in part.branch], color=figures_COLORS[k], marker=['o', 'x'][k], label=f'{temp} ℃', s=15)
            ax.set_xlabel(label)
        axes[0].legend(fontsize=6)
        save(fig, 'parameter_sensitivity')
    write_json(OUTPUT / 'figures' / 'figure_contract_and_QA.json', dict(backend='Python matplotlib', archetype='quantitative grid', conclusion='Show calibration limitations, inferred states, and conditional control outcomes without hiding failure', observations_per_temperature=184, exclusions=0, independent_replicates='not provided', uncertainty='no fabricated confidence intervals', calibration_runs=sources, exports=['SVG editable text', 'PDF embedded TrueType', 'PNG 300 dpi preview'], dimensions='3.6 to 9 inch width', language='中文', font=plt.rcParams['font.sans-serif'][0], visual_inspection='pending'))

# ==================== pipeline ====================
def selected_config():
    return Config(**json.loads((OUTPUT / 'selected_model.json').read_text(encoding='utf-8'))['config'])

def calibrate_all(workers=4, budget=16):
    tasks = [('B', 'H1', (-20, -25), 'P1', budget), ('A', 'H2', (-20, -25), 'P1', budget), ('A', 'H1', (-20, -25), 'P1', budget), ('B', 'H2', (-20, -25), 'P1', budget), ('B', 'H1', (-20, -25), 'P3', budget), ('B', 'H1', (-20,), 'P3', budget), ('B', 'H1', (-25,), 'P3', budget)]
    reports = []
    with ProcessPoolExecutor(max_workers=workers) as pool:
        for f in as_completed([pool.submit(calibrate_fit_job, *t) for t in tasks]):
            r = f.result()
            reports.append(r)
            print('calibrated', r['job'], flush=True)
    joint = next((r for r in reports if r['job'] == 'B_H1_P3_train-20_-25'))
    validation = [r for r in reports if len(r['train']) == 1]
    qualified = all((r['results'][str(t)].get('prediction_pass', False) for r in validation for t in (-20, -25) if t not in r['train']))
    write_json(OUTPUT / 'selected_model.json', dict(config=joint['config'], joint_fit=joint['job'], cross_temperature_qualified=qualified, validation=[r['job'] for r in validation], interpretation='conditional mechanism study' if not qualified else 'within-range validation'))
    return selected_config()

def replay_frozen():
    selected = json.loads((OUTPUT / 'selected_model.json').read_text(encoding='utf-8'))
    src = load_source()['experiments']
    for job in [selected['joint_fit'], *selected['validation']]:
        path = OUTPUT / 'calibration' / (job + '.json')
        report = json.loads(path.read_text(encoding='utf-8'))
        cfg = Config(**report['config'])
        for temp, d in src.items():
            r = experiment(cfg, temp, d, label='frozen_model')
            report['results'][str(temp)] = dict(role='fit' if temp in report['train'] else 'temperature_holdout', run=r['metrics']['candidate_id'], **waveform(r, d))
            export_mechanisms(r['metrics']['candidate_id'])
        report['frozen_replay'] = True
        write_json(path, report)
    return selected_config()

# ==================== optimize_voltage ====================
os.environ.setdefault('OMP_NUM_THREADS', '1')

os.environ.setdefault('OPENBLAS_NUM_THREADS', '1')

optimize_voltage_DEST = OUTPUT / 'voltage_optimization'

def optimize_voltage_fit_job(job, train, closure, budget=18, shape=False):
    data = load_source()['experiments']
    base = Config(tau_f=33.307669427458045, cl_proton_loss=closure != 'legacy', vapor_equilibrium='supercooled_liquid' if closure == 'liquid_cl' else 'ice_reference')
    x0 = [np.log(0.1723), np.log(1.54)]
    lo = [np.log(0.01), np.log(0.05)]
    hi = [np.log(3.0), np.log(40.0)]
    if closure in ('hydrated_cl', 'liquid_cl'):
        x0 += [3.0]
        lo += [0.0]
        hi += [6.0]
    calls = []
    cache = {}
    start = time.perf_counter()

    def unpack(x):
        gamma = float(x[2]) if len(x) > 2 else 0.0
        return base.changed(j0=float(np.exp(x[0]) * (14 / 3) ** gamma), tau_b=float(np.exp(x[1])), cathode_hydration_exponent=gamma)

    def residual(x):
        key = tuple(x)
        if key in cache:
            return cache[key]
        cfg = unpack(x)
        values = []
        scores = []
        for temp in train:
            d = data[temp]
            r = experiment(cfg, temp, d, label='vopt_' + job)
            complete = r['metrics']['stop_reason'] == 'horizon' and r['metrics']['numerically_accepted']
            v = np.interp(d.t, r['t'], r['V'][:, 0])
            tt = np.interp(d.t, r['t'], r['T'][:, 0])
            rr = np.r_[(v - d.V) / 0.03, (tt - d['T']) / 0.3]
            if shape:
                extra = np.sqrt(len(d) / 6) * np.array([v[0] - d.V.iloc[0], v.min() - d.V.min(), v[-1] - v.min() - (d.V.iloc[-1] - d.V.min())]) / 0.03
                rr = np.r_[rr, extra]
            if not complete:
                rr = np.full_like(rr, 100 + 10 * (d.t.iloc[-1] - r['t'][-1]))
            values.extend(rr)
            scores.append(dict(temp=temp, run=r['metrics']['candidate_id'], accepted=complete, **waveform(r, d)))
        cache[key] = np.asarray(values)
        entry = dict(call=len(calls) + 1, x=list(x), config=cfg.to_dict(), results=scores, elapsed_s=time.perf_counter() - start)
        calls.append(entry)
        write_json(optimize_voltage_DEST / (job + '_progress.json'), entry)
        print(job, len(calls), [(r['temp'], round(r.get('V_RMSE', 99), 4), round(r.get('T_RMSE', 99), 3)) for r in scores], flush=True)
        return cache[key]

    def jac(x):
        f0 = residual(x)
        cols = []
        for k in range(len(x)):
            xp = np.array(x, copy=True)
            step = 0.025 if x[k] + 0.025 < hi[k] else -0.025
            xp[k] += step
            cols.append((residual(xp) - f0) / step)
        return np.array(cols).T
    fit = least_squares(residual, x0, jac=jac, bounds=(lo, hi), max_nfev=budget, ftol=0.002, xtol=0.002, gtol=0.002, loss='soft_l1')
    cfg = unpack(fit.x)
    report = dict(job=job, train=train, closure=closure, config=cfg.to_dict(), nfev=fit.nfev, actual_parameter_evaluations=len(calls), cost=float(fit.cost), optimizer_message=fit.message, elapsed_s=time.perf_counter() - start, initial_parameters=x0, lower_bounds=lo, upper_bounds=hi, parameterization=['ln j0 effective at lambda=3', 'ln tau_b / s'] + (['hydration exponent'] if len(x0) > 2 else []), shape_features=shape, shape_feature_weight='sqrt(N/6)/0.03 for initial voltage, minimum voltage and recovery' if shape else None, fixed_tau_f_note='same previously calibrated freezing time in every variant; cross-temperature checks are retrospective', validation_scope='temperature excluded from objective; closure and freezing prior were developed after viewing both datasets', boundary_hit=bool(np.any(np.abs(fit.x - np.array(lo)) < 0.02) | np.any(np.abs(fit.x - np.array(hi)) < 0.02)), jacobian_singular_values=np.linalg.svd(fit.jac, compute_uv=False).tolist(), results={})
    for temp, d in data.items():
        r = experiment(cfg, temp, d, label='vopt_' + job)
        report['results'][str(temp)] = dict(role='training' if temp in train else 'retrospective_temperature_holdout', run=r['metrics']['candidate_id'], mass_residual=r['metrics']['mass_residual'], energy_residual=r['metrics']['energy_residual'], **waveform(r, d))
    write_json(optimize_voltage_DEST / (job + '.json'), report)
    write_json(optimize_voltage_DEST / (job + '_evaluations.json'), calls)
    return report

# ==================== verify_voltage ====================
os.environ.setdefault('OMP_NUM_THREADS', '1')

os.environ.setdefault('OPENBLAS_NUM_THREADS', '1')

verify_voltage_DEST = OUTPUT / 'voltage_optimization'

def readj(path):
    return json.loads(Path(path).read_text(encoding='utf-8'))

def verify_voltage_replay(job):
    name, config, temp = job
    d = load_source()['experiments'][temp]
    r = experiment(Config(**config), temp, d, label='vverify_' + name)
    row = dict(name=name, temp=temp, n_MEA=sum(config['mesh']), run=r['metrics']['candidate_id'], mass_residual=r['metrics']['mass_residual'], energy_residual=r['metrics']['energy_residual'], **waveform(r, d))
    write_json(verify_voltage_DEST / (name + f'_{temp}.json'), row)
    print('REPLAY', name, temp, row, flush=True)
    return row

def verify_voltage_verify(candidate):
    from scipy.optimize import brentq
    initial_rows = []
    for temp, d in load_source()['experiments'].items():

        def error(log_j0):
            m = Model(Config(T0=temp + 273.15, j0=float(np.exp(log_j0)), cl_proton_loss=True))
            return m.voltage(m.phases(m.initial()), float(d.j.iloc[0]))['V'][0] - float(d.V.iloc[0])
        initial_rows.append(dict(temperature_C=temp, initial_current_A_cm2=float(d.j.iloc[0]), initial_V_observed=float(d.V.iloc[0]), required_j0_effective_lambda3_A_m2=float(np.exp(brentq(error, -10, 10)))))
    write_json(verify_voltage_DEST / 'initial_voltage_constraint.json', dict(purpose='diagnostic only; no temperature-specific fitted parameters', rows=initial_rows))
    config = Config(**readj(verify_voltage_DEST / (candidate + '.json'))['config'])
    jobs = []
    for mesh, step, rtol in [((12, 6, 8, 10, 12), 0.2, 2e-06), ((24, 12, 16, 20, 24), 0.1, 1e-06)]:
        cfg = config.changed(mesh=mesh, max_step=step, rtol=rtol, board_n=4, end_n=8)
        jobs.extend([(candidate + '_grid' + str(sum(mesh)), cfg.to_dict(), temp) for temp in (-20, -25)])
    with ProcessPoolExecutor(4) as pool:
        rows = list(pool.map(verify_voltage_replay, jobs))
    checks = []
    for temp in (-20, -25):
        a, b = [load_run(next((x['run'] for x in rows if x['n_MEA'] == n and x['temp'] == temp))) for n in (48, 96)]
        d = load_source()['experiments'][temp]
        dv = float(np.max(abs(np.interp(d.t, a['t'], a['V'][:, 0]) - np.interp(d.t, b['t'], b['V'][:, 0]))))
        dt = float(np.max(abs(np.interp(d.t, a['t'], a['T'][:, 0]) - np.interp(d.t, b['t'], b['T'][:, 0]))))
        checks.append(dict(temp=temp, grid_pair=[48, 96], max_V_difference=dv, max_T_difference=dt, pass_=dv < 0.001 and dt < 0.02))
    old = readj(OUTPUT / 'selected_model.json')
    base = Config(**old['config'])
    r = experiment(base, -20, label='vverify_legacy')
    frozen = readj(OUTPUT / 'calibration' / (old['joint_fit'] + '.json'))
    saved = load_run(frozen['results']['-20']['run'])
    times = load_source()['experiments'][-20].t
    compat = {k: float(np.max(abs(np.interp(times, r['t'], r[k][:, 0]) - np.interp(times, saved['t'], saved[k][:, 0])))) for k in ('V', 'T', 'ice')}
    result = dict(candidate=candidate, replays=rows, grid_checks=checks, legacy_max_difference=compat)
    write_json(verify_voltage_DEST / (candidate + '_verification.json'), result)
    write_json(verify_voltage_DEST / 'verification.json', result)

def verify_voltage_report(candidate):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    configure_chinese_plotting(size=9)
    src = load_source()
    obs = src['experiments']
    specific = verify_voltage_DEST / (candidate + '_verification.json')
    validation = readj(specific if specific.exists() else verify_voltage_DEST / 'verification.json')
    assert validation['candidate'] == candidate
    new = readj(verify_voltage_DEST / (candidate + '.json'))
    oldsel = readj(OUTPUT / 'selected_model.json')
    old = readj(OUTPUT / 'calibration' / (oldsel['joint_fit'] + '.json'))
    family = candidate.split('_')[0]
    summary = []
    fine = {}
    lossrows = []
    tables = []
    for temp in (-20, -25):
        fine[temp] = next((x for x in validation['replays'] if x['n_MEA'] == 96 and x['temp'] == temp))
        for label, res in [('原模型联合校准 14格', old['results'][str(temp)]), ('改进模型联合校准 14格', new['results'][str(temp)]), ('改进模型复算 96格', fine[temp])]:
            summary.append(dict(model=label, temp_C=temp, **{k: res[k] for k in ['V_RMSE', 'T_RMSE', 'valley_time', 'depth_error', 'recovery_error', 'prediction_pass']}, run=res['run']))
        r = load_run(fine[temp]['run'])
        m = Model(Config(**r['config']))
        for i, (t, y, j) in enumerate(zip(r['t'], r['y'], r['j'])):
            s = m.phases(y)
            v = m.voltage(s, j)
            lossrows.append(dict(temp_C=temp, time_s=t, current_A_cm2=j, **{k: float(v[k][0]) for k in ['V', 'reversible', 'activation', 'ohmic', 'cl_ohmic', 'concentration', 'active_fraction']}, cathode_lambda_mean=float(np.average(s['lam'][0, m.g.clc], weights=m.g.dx[m.g.clc])), membrane_lambda_mean=r['lambda_mean'][i, 0]))
        original = src['tables'][0 if temp == -20 else 1]
        rows = [original[0]]
        for row in original[1:]:
            t = float(row[0])
            vm = float(np.interp(t, r['t'], r['V'][:, 0]))
            tm = float(np.interp(t, r['t'], r['T'][:, 0]))
            ice = float(np.interp(t, r['t'], r['ice'][:, 0]))
            rows.append([row[0], row[1], f'{vm:.4f}', f'{100 * abs(vm - float(row[1])) / abs(float(row[1])):.2f}', row[4], f'{tm:.2f}', f'{100 * abs(tm - float(row[4])) / abs(float(row[4])):.2f}', f'{ice:.6f}'])
        pd.DataFrame(rows[1:], columns=rows[0]).to_csv(verify_voltage_DEST / f'表{(1 if temp == -20 else 2)}_电压优化候选.csv', index=False, encoding='utf-8-sig')
        tables.append(rows)
    pd.DataFrame(summary).to_csv(verify_voltage_DEST / '误差对比.csv', index=False, encoding='utf-8-sig')
    pd.DataFrame(lossrows).to_csv(verify_voltage_DEST / '电压损失与含水量.csv', index=False, encoding='utf-8-sig')
    fig, axes = plt.subplots(3, 2, figsize=(9, 9), layout='constrained')
    curve_rows = []
    holdmetrics = []
    for k, temp in enumerate((-20, -25)):
        d = obs[temp]
        b = load_run(old['results'][str(temp)]['run'])
        n = load_run(fine[temp]['run'])
        ax = axes[0, k]
        ax.plot(d.t, d.V, color='black', label='实验数据')
        ax.plot(b['t'], b['V'][:, 0], color='#CC6677', label='原模型联合校准（14格）')
        ax.plot(n['t'], n['V'][:, 0], color='#0072B2', label='改进模型复算（96格）')
        ax.set(title=f'{temp} ℃ 联合校准', ylabel='电压（V）', xlabel='时间（s）')
        ax.legend(fontsize=7)
        train = -25 if temp == -20 else -20
        hj = readj(verify_voltage_DEST / (family + f'_train_m{abs(train)}.json'))
        h = load_run(hj['results'][str(temp)]['run'])
        oh = readj(OUTPUT / 'calibration' / f'B_H1_P3_train{train}.json')['results'][str(temp)]
        holdmetrics.append(dict(train=train, test=temp, old=oh, new=hj['results'][str(temp)], boundary_hit=hj['boundary_hit']))
        ax = axes[1, k]
        ax.plot(d.t, d.V, color='black', label='实验数据')
        ax.plot(h['t'], h['V'][:, 0], color='#009E73', label='回顾性留出检验（14格）')
        ax.set(title=f'训练：{train} ℃；检验：{temp} ℃', ylabel='电压（V）', xlabel='时间（s）')
        ax.legend(fontsize=7)
        ax = axes[2, k]
        ax.plot(d.t, d['T'], color='black', label='实验数据')
        ax.plot(b['t'], b['T'][:, 0], color='#CC6677', label='原模型')
        ax.plot(n['t'], n['T'][:, 0], color='#0072B2', label='改进模型')
        ax.set(ylabel='温度（℃）', xlabel='时间（s）')
        ax.legend(fontsize=7)
        for i, t in enumerate(d.t):
            curve_rows.append(dict(temp_C=temp, t_s=t, V_observed=d.V.iloc[i], V_original=float(np.interp(t, b['t'], b['V'][:, 0])), V_improved=float(np.interp(t, n['t'], n['V'][:, 0])), V_holdout=float(np.interp(t, h['t'], h['V'][:, 0])), T_observed=d['T'].iloc[i], T_improved=float(np.interp(t, n['t'], n['T'][:, 0]))))
    fig.savefig(verify_voltage_DEST / '电压预测优化对比.png', dpi=300)
    fig.savefig(verify_voltage_DEST / '电压预测优化对比.svg')
    fig.savefig(verify_voltage_DEST / '电压预测优化对比.pdf')
    plt.close(fig)
    pd.DataFrame(curve_rows).to_csv(verify_voltage_DEST / '逐点预测对比.csv', index=False, encoding='utf-8-sig')
    write_json(verify_voltage_DEST / '候选模型.json', dict(candidate=candidate, config=new['config'], fine_results=fine, retrospective_holdouts=holdmetrics, accepted_for_Q2_Q4=False, qualification='Not fully qualified; see unchanged waveform, thermal and grid gates', previous_control_model_preserved=oldsel['joint_fit']))
    lines = ['# 模型电压预测优化记录', '', '本轮对比催化层质子传导损失、含水活性闭合及过冷液水汽液平衡。以下数值全部来自重新积分的守恒模型，没有向电压输出添加按时间拟合的修正曲线。未改变原始实验数据、题设初温、初始膜含水量、接触电阻、活化能或膜扩散系数。', '', '## 预测改进与未通过项目', '', '| 模型 | 工况/℃ | 电压RMSE/V | 温度RMSE/℃ | 谷值时刻/s | 下陷幅度误差/V | 回升幅度误差/V | 全部预测门槛 |', '|---|---:|---:|---:|---:|---:|---:|---|']
    for s in summary:
        lines.append(f"| {s['model']} | {s['temp_C']} | {s['V_RMSE']:.5f} | {s['T_RMSE']:.4f} | {s['valley_time']:.1f} | {s['depth_error']:.5f} | {s['recovery_error']:.5f} | {('通过' if s['prediction_pass'] else '未通过')} |")
    lines += ['', '原门槛保持：电压RMSE≤0.03 V、温度RMSE≤0.3 ℃、谷值时刻误差≤1 s、下陷及回升幅度误差各≤0.03 V。改善电压不等于整套模型验收通过。', '', '在题给初态及活化能不变时，用扩展催化层电阻公式分别反解0.799 V首点，−20℃与−25℃要求的λ=3等效参考j0分别为0.216735和0.339039 A/m²，不能由同一个j0精确满足。这是当前共享参数模型与数据首点之间的约束冲突诊断，未据此增加两个温度专属修正量。复算见initial_voltage_constraint.json。', '', '## 温度交叉检查', '', '这些检查排除了目标温度的残差拟合，但模型结构和固定冻结参数来自此前已看过两工况的开发过程，属于回顾性温度留出检查，不是全新盲测。', '', '| 训练→检查/℃ | 原模型电压RMSE/V | 改进模型电压RMSE/V | 改进温度RMSE/℃ | 改进谷值/s | 参数贴边 |', '|---|---:|---:|---:|---:|---|']
    for h in holdmetrics:
        lines.append(f"| {h['train']} → {h['test']} | {h['old']['V_RMSE']:.5f} | {h['new']['V_RMSE']:.5f} | {h['new']['T_RMSE']:.4f} | {h['new']['valley_time']:.1f} | {('是' if h['boundary_hit'] else '否')} |")
    lines += ['', '## 方程与实现', '', '催化层额外质子损失使用耗散等效表达式：ΔV_CL = J∫(i_m/J)²/κ_eff dx，其中κ_eff=κ(T,λ)ω^1.5。题设均匀反应给出沿催化层线性变化的i_m/J；对每格的平方采用解析单元平均，均匀导电率时严格得到两层各L/(3κ_eff)，不随网格改变。它是对题目膜电阻简式的扩展，未再在热方程单独加一次焦耳热；总反应热仍为J(1.48−V)。', '', '含水活性采用可关闭的现象学候选：a_eff = ⟨(1−s_ice)^3.5 · min(λ_cCL/14,1)^γ⟩。γ由校准确定；λ/14幂律是本轮显式提出的候选闭合，不冒充题给公式或文献已验证定律。j0表示λ≥14时的参考交换电流密度；优化内部使用λ=3时等效j0，降低它与γ的参数相关。', '', '过冷液水分支使用液面Buck饱和压做汽液守恒分配，并保留有限速率液水冻结。冰库存不通过flash直接转为蒸气，没有补入无来源的升华速率。旧冰面参考分支作为对照保留；液水/冰/蒸气共存的非平衡过程仍是本构局限。', '', '## 参数与对照试验', '', f"候选编号：{candidate}。联合拟合参数贴边：{new['boundary_hit']}。", f"j0={new['config']['j0']:.8g} A/m²；τb={new['config']['tau_b']:.8g} s；γ={new['config']['cathode_hydration_exponent']:.8g}；τf固定为33.3076694 s。只校准三个参数，未把h=40、初温或热容作为补偿参数。", '', '拟合使用全部184点/工况的电压和温度残差，尺度分别为0.03 V和0.3 ℃，soft_l1损失。' + ('本候选还加入首点电压、全程最低电压和终点回升幅度三个残差，各乘sqrt(N/6)/0.03；只使用训练温度的观测。' if new.get('shape_features') else '') + '谷值、下陷和回升保留原门槛另行验收。每次评价有完整积分记录；有限次局部优化不构成全局最优证明。', '', '| 消融对照 | −20℃电压RMSE/V | −25℃电压RMSE/V |', '|---|---:|---:|']
    for file in ('cl_only_joint', 'hydrated_joint', 'liquid_joint', 'shape_joint'):
        if (verify_voltage_DEST / (file + '.json')).exists():
            rr = readj(verify_voltage_DEST / (file + '.json'))
            lines.append(f"| {file} | {rr['results']['-20']['V_RMSE']:.5f} | {rr['results']['-25']['V_RMSE']:.5f} |")
    lines += ['', '## 数值检查', '', '| 温度/℃ | 48→96格最大电压差/V | 最大温度差/℃ | 1 mV及0.02℃门槛 |', '|---|---:|---:|---|']
    for g in validation['grid_checks']:
        lines.append(f"| {g['temp']} | {g['max_V_difference']:.6f} | {g['max_T_difference']:.6f} | {('通过' if g['pass_'] else '未通过')} |")
    lines += ['', f"96格两工况最大水量残差={max((x['mass_residual'] for x in fine.values())):.3g}，最大能量残差={max((x['energy_residual'] for x in fine.values())):.3g}。", f"关闭新开关后，原−20℃轨迹回放最大差：{validation['legacy_max_difference']}。", '', '## 题目表1和表2的候选更新', '', '下表使用96格联合参数复算。保持题目原列、原实验值和8个取样时刻。它们与原冻结模型的表格分开保存，避免把新电压模型与旧Q2—Q4控制结果拼成同一套结果。相对误差由未舍入值计算。', '']
    for i, t in enumerate(tables):
        lines += ['### 表' + str(i + 1) + ' ' + str((-20, -25)[i]) + ' ℃', '', '| ' + ' | '.join(t[0]) + ' |', '| ' + ' | '.join(['---'] * 8) + ' |']
        lines += ['| ' + ' | '.join(row) + ' |' for row in t[1:]] + ['']
    lines += ['## 依据与复现', '', 'Edwards 与 Demuren的MEA模型支持将催化层损失、水吸脱附及导电性耦合；不将该常温稳态文献说成当前低温参数的验证。[原论文](https://link.springer.com/article/10.1007/s40095-018-0288-2)。', 'Yao等讨论过冷液水以及阳极失水和阴极堵塞的不同失效通路，支持保留相态及含水机制的对照；不据此认定当前实验谷值一定由结冰造成。[原论文](https://mdpi-res.com/d_attachment/energies/energies-13-00256/article_deploy/energies-13-00256-v2.pdf)。', '', '复现：在实施目录运行 `.venv/Scripts/python.exe q1.py optimize --budget 18 --workers 4`；再运行 `optimize_voltage.py --liquid --budget 18 --workers 3` 和 `optimize_voltage.py --shape --budget 12 --workers 3`；最后运行 `q1.py replay --candidate ' + candidate + '`。', '新增候选通过配置显式启用。outputs/selected_model.json及原Q2—Q4模型未替换；完整交付压缩包保留电压优化前的数值模型与控制结果，图中文字已统一为中文。本次源代码、候选配置、运行指标与表格以当前目录为准。']
    (verify_voltage_DEST / '电压预测优化报告.md').write_text('\n'.join(lines), encoding='utf-8')
    write_json(verify_voltage_DEST / 'result_manifest.json', dict(candidate=candidate, summary=summary, source_sha256=src['sha256'], reference_runs=[r['run'] for r in summary], retrospective_holdouts=holdmetrics, files=sorted((p.name for p in verify_voltage_DEST.iterdir() if p.is_file()))))

# ==================== analyze_ice ====================
matplotlib.use('Agg')

analyze_ice_DEST = OUTPUT / 'ice_analysis'

analyze_ice_DEST.mkdir(parents=True, exist_ok=True)

LAYER = ['阳极气体扩散层', '阳极催化层', '质子交换膜', '阴极催化层', '阴极气体扩散层']

analyze_ice_COLORS = ['#4477AA', '#EE6677', '#228833', '#AA3377', '#CCBB44']

def analyze_ice_read(p):
    return json.loads(p.read_text(encoding='utf-8'))

def analyze(name, temp, run):
    r = load_run(run)
    assert len(r['t']) == 184 and r['t'][-1] >= 36.6 - 1e-07
    assert r['metrics']['numerically_accepted']
    m = Model(Config(**r['config']))
    g = m.g
    rows = []
    fields = []
    s0 = m.phases(r['y'][0])
    initial = {key: float(np.sum(s0[key][0] * g.dx)) for key in ('mi', 'ml', 'mv', 'mb')}
    for k, (t, y, j) in enumerate(zip(r['t'], r['y'], r['j'])):
        s = m.phases(y)
        phi = s['mi'][0] / 920
        sl = s['ml'][0] / 990
        sat = np.divide(phi, g.eps, out=np.zeros_like(phi), where=g.porous)
        im = int(np.argmax(phi))
        ism = int(np.argmax(sat))
        cc = g.clc
        v = m.voltage(s, j)
        diagnostic = dict(s)
        diagnostic['mi'] = np.zeros_like(s['mi'])
        no_cover = m.voltage(diagnostic, j)
        mass = {key: float(np.sum(s[key][0] * g.dx)) for key in initial}
        charge = float(s['led'][-1])
        produced = charge * 10000.0 * MW / (2 * F)
        retained = sum(mass.values()) - sum(initial.values())
        out = produced - retained
        z = dict(model=name, temp_C=temp, run=run, time_s=float(t), phi_max=float(phi[im]), phi_layer=LAYER[g.layer[im]], phi_x_um=float(g.x[im] * 1000000.0), saturation_max=float(sat[ism]), saturation_layer=LAYER[g.layer[ism]], saturation_x_um=float(g.x[ism] * 1000000.0), phi_cCL_max=float(phi[cc].max()), phi_MEA_mean=float(np.average(phi, weights=g.dx)), saturation_at_phi_peak=float(sat[im]), gas_pore_min=float(s['eg'][0, g.porous].min()), occupied_pore_max=float(np.max((phi[g.porous] + sl[g.porous]) / g.eps[g.porous])), T_max_C=float(s['local_T'].max() - TM), T_mean_C=float(r['T'][k, 0]), V=float(v['V'][0]), ice_coverage_voltage_loss_mV=float((no_cover['V'][0] - v['V'][0]) * 1000), freeze_rate_kg_m2_s=float(np.sum(s['ml'][0] * np.maximum(TM - s['local_T'][0], 0) / (20 * m.cfg.tau_f) * g.dx)), melt_rate_kg_m2_s=float(np.sum(np.maximum(s['mi'][0], 0) * np.maximum(s['local_T'][0] - TM, 0) / (5 * m.cfg.tau_m) * g.dx)), produced_kg_m2=produced, outward_kg_m2=out, bound_delta_kg_m2=mass['mb'] - initial['mb'], **{key + '_kg_m2': mass[key] for key in mass}, active_fraction=float(v['active_fraction'][0]), lambda_cCL=float(np.average(s['lam'][0, cc], weights=g.dx[cc])), mass_residual=float(r['metrics']['mass_residual']), **{key + '_V': float(v[key][0]) for key in ('activation', 'ohmic', 'cl_ohmic', 'concentration', 'reversible')})
        for layer in range(5):
            z[f'phi_layer{layer}_max'] = float(phi[g.layer == layer].max())
        rows.append(z)
        if name == '改进模型96格':
            for i in range(g.n):
                fields.append(dict(temp_C=temp, time_s=t, x_um=g.x[i] * 1000000.0, layer=LAYER[g.layer[i]], phi=phi[i], saturation=sat[i], liquid_bulk=sl[i], gas_pore=s['eg'][0, i], T_C=s['local_T'][0, i] - TM))
    d = pd.DataFrame(rows)
    assert np.allclose(d.phi_max, r['ice'][:, 0], atol=1e-12)
    assert np.allclose(d.saturation_max, r['ice_saturation'][:, 0], atol=1e-12)
    peak = d.loc[d.phi_max.idxmax()].to_dict()
    valley = d.loc[d.V.idxmin()].to_dict()
    onset = d[d.phi_max >= 1e-06]
    return (d, fields, dict(model=name, temp_C=temp, run=run, n_MEA=g.n, peak=peak, valley=valley, onset_threshold=1e-06, onset_time_s=None if onset.empty else float(onset.time_s.iloc[0]), initial_inventory=initial, peak_grid_sampled=True))

def analyze_ice_main():
    configure_chinese_plotting(9)
    candidate = analyze_ice_read(OUTPUT / 'voltage_optimization/候选模型.json')
    new = analyze_ice_read(OUTPUT / 'voltage_optimization/shape_joint.json')
    verify = analyze_ice_read(OUTPUT / 'voltage_optimization/shape_joint_verification.json')
    old = analyze_ice_read(OUTPUT / 'selected_model.json')
    oldfit = analyze_ice_read(OUTPUT / 'calibration' / (old['joint_fit'] + '.json'))
    curves = []
    fieldrows = []
    summary = []
    for temp in (-20, -25):
        jobs = [('原模型14格', oldfit['results'][str(temp)]['run']), ('改进模型14格', new['results'][str(temp)]['run'])]
        jobs += [(f'改进模型{n}格', next((x['run'] for x in verify['replays'] if x['temp'] == temp and x['n_MEA'] == n))) for n in (48, 96)]
        extra = analyze_ice_DEST / f'grid192_{temp}.json'
        if extra.exists():
            jobs.append(('改进模型192格', analyze_ice_read(extra)['run']))
        for name, run in jobs:
            d, fields, s = analyze(name, temp, run)
            curves.append(d)
            fieldrows += fields
            summary.append(s)
    allcurves = pd.concat(curves, ignore_index=True)
    allcurves.to_csv(analyze_ice_DEST / '冰分数与水量逐时诊断.csv', index=False, encoding='utf-8-sig')
    pd.DataFrame(fieldrows).to_csv(analyze_ice_DEST / '96格冰水空间分布.csv', index=False, encoding='utf-8-sig')
    write_json(analyze_ice_DEST / 'analysis_summary.json', summary)
    peaks = pd.DataFrame([dict(n_MEA=x['n_MEA'], onset_time_s=x['onset_time_s'], **x['peak']) for x in summary])
    peaks.to_csv(analyze_ice_DEST / '峰值与网格对照.csv', index=False, encoding='utf-8-sig')
    table = []
    for temp in (-20, -25):
        d = allcurves[(allcurves.model == '改进模型96格') & (allcurves.temp_C == temp)]
        for row in load_source()['tables'][0 if temp == -20 else 1][1:]:
            t = float(row[0])
            rec = d.iloc[int(np.argmin(abs(d.time_s - t)))]
            assert abs(rec.time_s - t) < 1e-08
            table.append({'初温/℃': temp, '时间/s': t, '最大冰体积分数': rec.phi_max, '最大冰体积分数/%': rec.phi_max * 100, '最大孔隙冰饱和度/%': rec.saturation_max * 100, '峰值所在层': rec.phi_layer if rec.phi_max >= 1e-06 else '低于诊断阈值', '阴极催化层最大冰体积分数': rec.phi_cCL_max})
    pd.DataFrame(table).to_csv(analyze_ice_DEST / '题目时刻冰分数补充表.csv', index=False, encoding='utf-8-sig')

    def save(fig, name):
        fig.savefig(analyze_ice_DEST / (name + '.png'), dpi=300, bbox_inches='tight')
        fig.savefig(analyze_ice_DEST / (name + '.svg'), bbox_inches='tight')
        fig.savefig(analyze_ice_DEST / (name + '.pdf'), bbox_inches='tight')
        plt.close(fig)
    fig, axes = plt.subplots(3, 2, figsize=(8.6, 8.2), layout='constrained')
    for k, temp in enumerate((-20, -25)):
        d = allcurves[(allcurves.model == '改进模型96格') & (allcurves.temp_C == temp)]
        b = allcurves[(allcurves.model == '原模型14格') & (allcurves.temp_C == temp)]
        axes[0, k].plot(b.time_s, b.phi_max * 100, label='原模型（14格）', color=analyze_ice_COLORS[1])
        axes[0, k].plot(d.time_s, d.phi_max * 100, label='改进模型（96格）', color=analyze_ice_COLORS[0])
        axes[0, k].set(title=f'{temp} ℃ 工况', ylabel='最大冰体积分数（%）')
        axes[0, k].legend(fontsize=7)
        axes[1, k].plot(d.time_s, d.saturation_max * 100, label='最大孔隙冰饱和度', color=analyze_ice_COLORS[3])
        axes[1, k].plot(d.time_s, d.occupied_pore_max * 100, label='最大冰液合计占孔率', color=analyze_ice_COLORS[2])
        axes[1, k].set_ylabel('占初始孔隙体积比例（%）')
        axes[1, k].legend(fontsize=7)
        axes[2, k].plot(d.time_s, d.ice_coverage_voltage_loss_mV, color=analyze_ice_COLORS[0])
        valley = d.loc[d.V.idxmin()]
        axes[2, k].axvline(valley.time_s, color='#777777', ls='--', label='预测电压谷值时刻')
        axes[2, k].set_ylabel('冰覆盖直接电压损失（mV）')
        axes[2, k].legend(fontsize=7)
        for ax in axes[:, k]:
            ax.set_xlabel('时间（s）')
    fig.suptitle('冰体积分数、占孔率与电压影响（模型推断）')
    save(fig, '冰分数与电压影响分析')
    fig, axes = plt.subplots(2, 2, figsize=(9, 6.4), layout='constrained')
    f = pd.DataFrame(fieldrows)
    for k, temp in enumerate((-20, -25)):
        part = f[f.temp_C == temp]
        end = part[part.time_s == part.time_s.max()]
        axes[0, k].plot(end.x_um, end.phi * 100, color=analyze_ice_COLORS[0], label='总体积冰分数')
        axes[0, k].plot(end.x_um, end.saturation * 100, color=analyze_ice_COLORS[3], label='孔隙冰饱和度')
        for x in np.cumsum(THICKNESS)[:-1] * 1000000.0:
            axes[0, k].axvline(x, color='#aaaaaa', lw=0.6, ls=':')
        axes[0, k].set(title=f'{temp} ℃，36.6 s空间分布', xlabel='膜电极内位置（μm）', ylabel='体积比例（%）')
        axes[0, k].legend(fontsize=7)
        d = allcurves[(allcurves.model == '改进模型96格') & (allcurves.temp_C == temp)]
        for field, label, color in [('produced_kg_m2', '累计产水', analyze_ice_COLORS[0]), ('outward_kg_m2', '累计净排出', analyze_ice_COLORS[1]), ('bound_delta_kg_m2', '束缚水增量', analyze_ice_COLORS[2]), ('mi_kg_m2', '冰库存', analyze_ice_COLORS[3]), ('ml_kg_m2', '液水库存', analyze_ice_COLORS[4])]:
            axes[1, k].plot(d.time_s, d[field] * 1000, label=label, color=color)
        axes[1, k].set(xlabel='时间（s）', ylabel='单位面积水量（g/m²）')
        axes[1, k].legend(fontsize=7, ncol=2)
    fig.suptitle('末时刻空间分布与水量去向（模型推断）')
    save(fig, '冰分布与水量去向')
    fig, axes = plt.subplots(1, 2, figsize=(8.6, 3), layout='constrained')
    for k, temp in enumerate((-20, -25)):
        part = peaks[(peaks.temp_C == temp) & peaks.model.str.startswith('改进')].sort_values('n_MEA')
        axes[0].plot(part.n_MEA, part.phi_max * 100, 'o-', color=analyze_ice_COLORS[k], label=f'{temp} ℃')
        valid = part[part.onset_time_s.notna()]
        axes[1].plot(valid.n_MEA, valid.onset_time_s, 'o-', color=analyze_ice_COLORS[k], label=f'{temp} ℃')
    axes[0].set(xlabel='膜电极网格数', ylabel='全过程最大冰体积分数（%）')
    axes[1].set(xlabel='膜电极网格数', ylabel='首次超过诊断阈值的时刻（s）', title='−20 ℃的14格未达到阈值')
    for ax in axes:
        ax.set_xticks([14, 48, 96, 192])
        ax.legend(fontsize=7)
    fig.suptitle('冰峰值与起冰时刻的网格敏感性（阈值为0.000001）')
    save(fig, '冰峰值网格敏感性')
    print(peaks[['model', 'temp_C', 'phi_max', 'saturation_max', 'phi_layer', 'phi_x_um', 'mi_kg_m2', 'outward_kg_m2', 'produced_kg_m2', 'bound_delta_kg_m2', 'ice_coverage_voltage_loss_mV']].to_string(index=False))
    print('VALLEYS', [(s['temp_C'], s['valley']['time_s'], s['valley']['phi_max'], s['onset_time_s']) for s in summary if s['n_MEA'] == 96])

# ==================== refine_ice ====================
os.environ.setdefault('OMP_NUM_THREADS', '1')

os.environ.setdefault('OPENBLAS_NUM_THREADS', '1')

def refine_ice_run(temp):
    p = analyze_ice_read(OUTPUT / 'voltage_optimization/候选模型.json')
    base = Config(**load_run(p['fine_results'][str(temp)]['run'])['config'])
    cfg = base.changed(mesh=(48, 24, 32, 40, 48))
    result = experiment(cfg, temp, label='ice_grid192')
    row = dict(temp=temp, n_MEA=192, run=result['metrics']['candidate_id'], metrics=result['metrics'])
    write_json(OUTPUT / 'ice_analysis' / f'grid192_{temp}.json', row)
    print(temp, row['run'], row['metrics']['ice_max'], flush=True)
    return row

# ==================== report_ice ====================
D = OUTPUT / 'ice_analysis'

def report_ice_report():
    data = pd.read_csv(D / '冰分数与水量逐时诊断.csv')
    s = json.loads((D / 'analysis_summary.json').read_text(encoding='utf-8'))
    main = {t: next((x for x in s if x['temp_C'] == t and x['model'] == '改进模型96格')) for t in (-20, -25)}
    lines = ['# 模型最大冰体积分数专项分析', '', '分析对象为电压优化候选shape_joint；96格轨迹是当前候选表1、表2的来源。另用原模型14格、改进模型14/48格，以及本轮192格复算检查数值稳定性。全部数值由状态轨迹重新提取；未重新拟合参数、修改实验数据或替换冻结的Q2—Q4模型。', '', '## 1. 首要结论', '', '当前模型的电压先下降后回升，不能解释为“先结冰堵塞、随后融冰恢复”。预测谷值附近的冰体积分数仅约10⁻¹⁹，属于数值舍入量级；可辨识积冰出现在更晚时段。96格两工况所有膜电极网格在36.6 s内均低于0 ℃，模型融化源项始终为零。该结论针对当前模型机制，不等于已经证明真实实验没有早期结冰。', '', '## 2. 三个统计量必须区分', '', '- 局部总体积冰分数：φᵢ(x,t)=mᵢ(x,t)/920；mᵢ以控制体总体积为基准，单位kg/m³。题目所需最大冰体积分数是maxₖ,ₓ φᵢ(x,t)。', '- 孔隙冰饱和度：sᵢ(x,t)=φᵢ(x,t)/ε₀(x)。它与总体积冰分数分母不同。', '- 全过程峰值：maxₜ maxₖ,ₓ φᵢ，与表格每个指定时刻的空间最大值不同。本文峰值来自0.2 s输出采样；本轮192格保留相同采样间隔。', '', '阴极催化层孔隙率为0.4207，阴极气体扩散层为0.8。因为孔隙率不同，全域最大φᵢ和最大sᵢ可能不在同一个网格，不能用全域max φᵢ除以某一个固定孔隙率得到全域max sᵢ。', '', '## 3. 当前96格候选的数值', '', '| 初温/℃ | 全过程最大φᵢ | 换算百分比 | 时刻/s | 所在层 | 位置/μm | 最大sᵢ/% |', '|---|---:|---:|---:|---|---:|---:|']
    for t, x in main.items():
        p = x['peak']
        lines.append(f"| {t} | {p['phi_max']:.8f} | {p['phi_max'] * 100:.5f}% | {p['time_s']:.1f} | {p['phi_layer']} | {p['phi_x_um']:.3f} | {p['saturation_max'] * 100:.5f}% |")
    lines += ['', '位置以阳极气体扩散层外侧为原点；各层依次为阳极扩散层0–150、阳极催化层150–153.4、膜153.4–165.4、阴极催化层165.4–176.7、阴极扩散层176.7–326.7 μm。网格中心坐标并不代表连续介质中已精确定位的峰值。', '', '−25 ℃时最大φᵢ出现在阴极扩散层靠近催化层一侧，最大sᵢ则出现在阴极催化层。两者都很小，但不能由此推断模型已经通过实验验证。', '', '| 初温/℃ | 首次φᵢ≥10⁻⁶的输出时刻/s | 电压谷值时刻/s | 谷值时刻φᵢ | 最大冰液合计占孔率/% |', '|---|---:|---:|---:|---:|']
    for t, x in main.items():
        d = data[(data.model == '改进模型96格') & (data.temp_C == t)]
        v = x['valley']
        lines.append(f"| {t} | {x['onset_time_s']:.1f} | {v['time_s']:.1f} | {v['phi_max']:.3e} | {d.occupied_pore_max.max() * 100:.4f} |")
    lines += ['', '10⁻⁶仅为本次报告区分舍入噪声与积冰的诊断阈值，不是成核阈值、实验检出限或题设约束。起冰时刻是首次跨过该阈值的0.2 s输出时刻。初态无冰；显示为0.000000不等于所有后续时刻都严格无冰。', '', '## 4. 水为什么没有全部冻在催化层', '', '当前闭合允许水在催化层与扩散层间迁移，也允许催化层离聚物吸水和外边界汽相排出。outlet=closed只关闭液态水外排，未关闭水蒸气外排。因此水可先被离聚物吸收或经蒸气排出，剩余液水才通过有限速率源项冻结。', '', '以下为36.6 s单位活性面积水量，单位g/m²。排出量由累计产水减总水库存增量得到；束缚水列必须使用相对初态的增量，不能把初始膜内水再算作产水。初始自由水与冰均为零。', '', '| 初温/℃ | 累计产水 | 累计净排出 | 束缚水增量 | 液水库存 | 冰库存 | 蒸气库存 |', '|---|---:|---:|---:|---:|---:|---:|']
    for t, x in main.items():
        end = data[(data.model == '改进模型96格') & (data.temp_C == t)].iloc[-1]
        keys = ['produced_kg_m2', 'outward_kg_m2', 'bound_delta_kg_m2', 'ml_kg_m2', 'mi_kg_m2', 'mv_kg_m2']
        lines.append('| ' + str(t) + ' | ' + ' | '.join((f'{end[k] * 1000:.6f}' for k in keys)) + ' |')
    lines += ['', '净排出可含初态束缚水的解吸贡献，不能把它解释为对新生成水分子的溯源比例。当前边界为干气，排出量依赖该边界假设，不能直接当成实测排水能力。水量守恒残差通过仅说明数值收支自洽，不能验证相分配闭合。', '', '冻结源项为r_f=m_l·max(273.15−T,0)/(20τ_f)，本候选τ_f=33.3076694 s沿用前轮结果，电压优化时未重估；在−20 ℃和−25 ℃恒温、孤立液水近似下，局部冻结时间常数分别为33.31 s和26.65 s，温度升高时冻结进一步变慢。这是现象学闭合，不含显式随机成核或实验测得的诱导期。', '', '## 5. 冰对电压的作用有多大', '', '固定每一时刻的温度、含水量、氧浓度与气孔隙率，只在活化公式中去掉(1−sᵢ)^3.5冰覆盖因子，得到如下直接压降。该诊断不是完整“无冰”反事实积分，未计入冰改变传质、热量与水分布的间接作用。', '', '| 初温/℃ | 36.6 s冰覆盖直接压降/mV | 谷值到终点电压回升/mV | 谷值到终点活化损失下降/mV | 谷值到终点欧姆损失下降/mV |', '|---|---:|---:|---:|---:|']
    for t, x in main.items():
        d = data[(data.model == '改进模型96格') & (data.temp_C == t)]
        v = d.loc[d.V.idxmin()]
        e = d.iloc[-1]
        lines.append(f'| {t} | {e.ice_coverage_voltage_loss_mV:.5f} | {(e.V - v.V) * 1000:.3f} | {(v.activation_V - e.activation_V) * 1000:.3f} | {(v.ohmic_V - e.ohmic_V) * 1000:.3f} |')
    lines += ['', '原模型与改进模型曲线同时存在结构和网格差异，不能将冰量变化归因于某个单独改动。当前模型将回升主要归于含水状态和温度改善后活化、欧姆损失下降；电流在增加，而冰在末段才开始积累，且未融化。各项变化是模型内部的损失分解，不是独立实验证明的因果贡献。', '', '## 6. 网格敏感性：最大值仍未收敛', '', '| 初温/℃ | 膜电极网格数 | 全过程最大φᵢ | 峰值所在层 | 首次跨10⁻⁶/s |', '|---|---:|---:|---|---:|']
    for t in (-20, -25):
        for x in s:
            if x['temp_C'] != t or not x['model'].startswith('改进'):
                continue
            p = x['peak']
            on = '未达到' if x['onset_time_s'] is None else f"{x['onset_time_s']:.1f}"
            layer = p['phi_layer'] if p['phi_max'] >= 1e-06 else '舍入量级，无可辨识峰值'
            lines.append(f"| {t} | {x['n_MEA']} | {p['phi_max']:.8g} | {layer} | {on} |")
    lines += ['', '48→96格改变了空间网格与时间步/容差设置，因此属于综合分辨率检查。本轮96→192格仅加密膜电极网格，板件网格、rtol=10⁻⁶、max_step=0.1 s和参数保持96格设置，用于更干净地检查空间离散误差。', '']
    for t in (-20, -25):
        a = main[t]['peak']['phi_max']
        extra = next((x for x in s if x['temp_C'] == t and x['n_MEA'] == 192), None)
        if extra:
            b = extra['peak']['phi_max']
            lines.append(f'- {t} ℃：96→192格峰值由{a:.8g}变为{b:.8g}，相对96格变化{(b / a - 1) * 100:.2f}%；绝对差{abs(b - a):.6g}，折合{abs(b - a) * 100:.6g}个百分点。')
    lines += ['', '局部最大值对窄积水区域较敏感；不能把96格的小数位当成有效精度，也不能从守恒残差很小推出局部峰值已收敛。本次未预设或放宽冰峰值验收阈值，当前精确数值仍按条件性网格结果报告。', '', '## 7. 题设约束与下一阶段', '', '题设0.99是总体积冰分数阈值；本题干孔隙率最高仅0.8，物理可行状态在达到0.99前已耗尽孔容。因此不能用“φᵢ<0.99”单独证明无堵塞或启动成功。应同时报告最大sᵢ、最大冰液合计占孔率、最小气孔隙率、供氧约束和全过程最低电压。', '', '现有附件未提供直接冰量观测，τ_f也未在本轮电压优化中重新识别。后续应先给冰峰值、冰库存和起冰时间单列网格验收标准，再在固定网格下做冻结/吸附时间常数及边界湿度敏感性；若要声称冰量可信，需要独立含冰量或相态证据。当前不宜为了得到更大的冰曲线而任意调快冻结速率，也不宜因新电压拟合改善就替换Q2—Q4的模型。', '', '![冰峰值网格敏感性](冰峰值网格敏感性.png)', '', '## 8. 文件与复现', '', '运行 `python q1.py ice-refine` 生成192格核查轨迹；随后运行 `python q1.py ice-analysis` 和 `python q1.py ice-analysis`。', '诊断源数据：冰分数与水量逐时诊断.csv、96格冰水空间分布.csv、峰值与网格对照.csv；题设8个时刻的补充量见题目时刻冰分数补充表.csv。所有图中文字为中文。', '', '![冰分数与电压影响](冰分数与电压影响分析.png)', '', '![冰分布与水量去向](冰分布与水量去向.png)', '']
    (D / '最大冰体积分数分析.md').write_text('\n'.join(lines), encoding='utf-8')

# ==================== fit_ice_curve ====================
matplotlib.use('Agg')

fit_ice_curve_ROOT = Path(__file__).resolve().parent

fit_ice_curve_DEST = _project_storage() / 'outputs/ice_analysis'

fit_ice_curve_SOURCE = fit_ice_curve_DEST / '冰分数与水量逐时诊断.csv'

def fit_ice_curve_main():
    configure_chinese_plotting(10)
    data = pd.read_csv(fit_ice_curve_SOURCE)
    fig, axes = plt.subplots(1, 2, figsize=(9, 3.8), layout='constrained')
    results = []
    for ax, temp in zip(axes, (-20, -25)):
        d = data[(data.temp_C == temp) & (data.model == '改进模型192格')].sort_values('time_s')
        coarse = data[(data.temp_C == temp) & (data.model == '改进模型96格')].sort_values('time_s')
        t = d.time_s.to_numpy()
        y = d.phi_max.to_numpy()
        end = float(t[-1])
        scale = float(y.max())
        assert len(t) == 184 and np.isfinite(y).all()

        def predict(x, tt):
            amplitude, onset, power = x
            return amplitude * np.maximum((tt - onset) / (end - onset), 0) ** power
        starts = []
        onset = float(d[d.phi_max >= 1e-06].time_s.iloc[0])
        for power in (1.0, 2.0, 3.0):
            result = least_squares(lambda x: (predict(x, t) - y) / scale, [scale, onset - 0.2, power], bounds=([0, 0, 1], [3 * scale, end - 0.01, 5]), x_scale=[scale, 10, 2], ftol=1e-12, xtol=1e-12, gtol=1e-12, max_nfev=1000)
            starts.append(result)
        fit = min(starts, key=lambda r: np.sum(r.fun ** 2))
        pred = predict(fit.x, t)
        active = y >= 1e-06
        summary = {'temp_C': temp, 'source_model': '改进模型192格', 'n': len(t), 'formula': 'phi(t)=A*[max((t-t0)/(36.6-t0),0)]^p', 'domain_s': [0, end], 'A': float(fit.x[0]), 't0_s': float(fit.x[1]), 'p': float(fit.x[2]), 'RMSE': float(np.sqrt(np.mean((pred - y) ** 2))), 'active_RMSE': float(np.sqrt(np.mean((pred[active] - y[active]) ** 2))), 'active_n': int(active.sum()), 'max_abs_error': float(np.max(abs(pred - y))), 'R2': float(1 - np.sum((pred - y) ** 2) / np.sum((y - y.mean()) ** 2)), 'model_threshold_crossing_s': onset, 'optimizer_success': bool(fit.success), 'warning': '仅拟合模型输出，非实验冰量验证；拟合t0不等于真实成核时刻；不用于区间外外推。'}
        results.append(summary)
        fine = np.linspace(0, end, 1200)
        ax.plot(coarse.time_s, coarse.phi_max * 100, ls='--', lw=1.25, color='#999999', label='96格模型计算')
        ax.scatter(t, y * 100, s=8, color='#333333', alpha=0.6, label='192格模型计算', zorder=3)
        ax.plot(fine, predict(fit.x, fine) * 100, lw=2, color='#0072B2', label='分段幂函数拟合', zorder=4)
        ax.axvline(onset, color='#777777', lw=0.7, ls=':')
        ax.annotate(f'模型首次跨阈值：{onset:.1f} s', xy=(onset, 0), xycoords='data', xytext=(0.06, 0.44), textcoords='axes fraction', fontsize=8, arrowprops={'arrowstyle': '->', 'color': '#777777', 'lw': 0.7})
        ax.set(title=f'{temp} ℃ 工况', xlabel='时间（s）', ylabel='最大冰体积分数（%）', xlim=(0, 37.5), ylim=(-scale * 3, scale * 118))
        ax.legend(loc='upper left', fontsize=8)
        ax.text(0.06, 0.64, f"拟合决定系数：{summary['R2']:.5f}", transform=ax.transAxes, fontsize=8)
    fig.suptitle('最大冰体积分数拟合曲线', fontsize=13)
    fig.supxlabel('基于模型输出拟合；诊断阈值为0.000001；未作为实测冰量验证', fontsize=8)
    for ext in ('png', 'svg', 'pdf'):
        fig.savefig(fit_ice_curve_DEST / ('最大冰体积分数拟合曲线.' + ext), dpi=300, bbox_inches='tight')
    plt.close(fig)
    out = {'source_sha256': hashlib.sha256(fit_ice_curve_SOURCE.read_bytes()).hexdigest(), 'fit_type': '分段幂函数最小二乘拟合', 'fits': results}
    (fit_ice_curve_DEST / '冰分数拟合参数.json').write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding='utf-8')
    lines = ['# 最大冰体积分数拟合曲线', '', '拟合对象：两个温度下192格模型的全部184个输出点。96格结果仅作数值敏感性对照，不参与此次拟合。冰量没有实验观测，本拟合只提供轨迹的简洁近似，不是模型验证。', '', '函数：φ(t)=A[max((t−t₀)/(36.6−t₀),0)]ᵖ，适用范围0≤t≤36.6 s；φ为无量纲总体积冰分数，图中纵轴为100φ（%）。', '', '| 初温/℃ | A | t₀/s | p | 全程RMSE | 积冰段RMSE | 决定系数 |', '|---|---:|---:|---:|---:|---:|---:|']
    for r in results:
        lines.append(f"| {r['temp_C']} | {r['A']:.9g} | {r['t0_s']:.6f} | {r['p']:.6f} | {r['RMSE']:.6g} | {r['active_RMSE']:.6g} | {r['R2']:.6f} |")
    lines += ['', 'RMSE以无量纲冰分数计；积冰段指模型φ≥10⁻⁶的采样点。拟合起点t₀是经验参数，不等于模型首次跨诊断阈值时刻，也不等于真实成核时间。全程决定系数包含长时间近零段；不能将高决定系数解读为含冰量实验验证通过。禁止用该幂函数替代守恒模型进行时间区间外或其他工况的预测。', '', '图中每一个192格输出点均保留；没有修改原轨迹。相邻网格峰值差异仍存在。', '', '![最大冰体积分数拟合曲线](最大冰体积分数拟合曲线.png)']
    (fit_ice_curve_DEST / '冰分数拟合说明.md').write_text('\n'.join(lines) + '\n', encoding='utf-8')
    print(json.dumps(results, ensure_ascii=False, indent=2))

# ==================== plot_voltage_fit ====================
matplotlib.use('Agg')

def plot_voltage_fit_main():
    configure_chinese_plotting(10)
    source = load_source()
    candidate = json.loads((OUTPUT / 'voltage_optimization/候选模型.json').read_text(encoding='utf-8'))
    dest = OUTPUT / 'voltage_optimization'
    fig, axes = plt.subplots(1, 2, figsize=(9, 3.8), layout='constrained')
    results = []
    samples = []
    for ax, temp in zip(axes, (-20, -25)):
        d = source['experiments'][temp]
        fine = json.loads((OUTPUT / f'ice_analysis/grid192_{temp}.json').read_text(encoding='utf-8'))
        runs = {96: candidate['fine_results'][str(temp)]['run'], 192: fine['run']}
        ax.plot(d.t, d.V, color='#333333', lw=1.5, label='实验电压（184点）', zorder=4)
        ax.scatter(d.t, d.V, s=4, color='#333333', zorder=4)
        predictions = {}
        rmse = {}
        for n, color, style, label in [(96, '#D98B32', '--', '当前模型（96格）'), (192, '#0072B2', '-', '同参数复算（192格）')]:
            r = load_run(runs[n])
            assert len(r['t']) == 184 and r['t'][-1] >= d.t.iloc[-1] - 1e-07
            y = np.interp(d.t, r['t'], r['V'][:, 0])
            predictions[n] = y
            rmse[n] = float(np.sqrt(np.mean((y - d.V.to_numpy()) ** 2)))
            results.append(dict(temp_C=temp, n_MEA=n, run=runs[n], n_observations=len(d), V_RMSE=rmse[n], V_MAE=float(np.mean(abs(y - d.V.to_numpy()))), V_min=float(y.min()), valley_time_s=float(d.t.iloc[np.argmin(y)])))
            ax.plot(r['t'], r['V'][:, 0], color=color, ls=style, lw=1.6, label=label, zorder=3)
        for i, t in enumerate(d.t):
            samples.append(dict(temp_C=temp, t_s=float(t), V_observed=float(d.V.iloc[i]), V_96=float(predictions[96][i]), V_192=float(predictions[192][i])))
        ax.set(title=f'{temp} ℃ 工况', xlabel='时间（s）', ylabel='单片电压（V）', xlim=(0, 37.5), ylim=(0.46, 0.83))
        ax.legend(loc='upper right', fontsize=8)
        ax.text(0.035, 0.06, f'均方根误差：\n96格：{rmse[96] * 1000:.2f} mV\n192格：{rmse[192] * 1000:.2f} mV', transform=ax.transAxes, fontsize=8, va='bottom')
    fig.suptitle('电压拟合曲线：实验与模型对比', fontsize=13)
    fig.supxlabel('联合校准参数下的回代与网格复算；未另加经验修正曲线', fontsize=8)
    fig.savefig(dest / '电压拟合曲线.png', dpi=300, bbox_inches='tight')
    fig.savefig(dest / '电压拟合曲线.svg', bbox_inches='tight')
    fig.savefig(dest / '电压拟合曲线.pdf', bbox_inches='tight')
    plt.close(fig)
    write_json(dest / '电压拟合曲线数据.json', dict(source_sha256=source['sha256'], candidate=candidate['candidate'], interpretation='联合参数校准回代及同参数192格复算，不是独立盲测；全部184点/工况保留，无后处理平滑或附加经验拟合。', metrics=results, samples=samples))
    print(json.dumps(results, ensure_ascii=False, indent=2))

# ==================== review_figures ====================
def contact_sheets(include_voltage=True, dest=None):
    files = sorted((OUTPUT / 'figures').glob('*.png'))
    voltage = OUTPUT / 'voltage_optimization/电压预测优化对比.png'
    if include_voltage and voltage.exists():
        files.append(voltage)
    dest = dest or OUTPUT / 'qa'
    dest.mkdir(exist_ok=True, parents=True)
    font = ImageFont.truetype(configure_chinese_plotting(), 24)
    manifest = []
    for first in range(0, len(files), 4):
        canvas = Image.new('RGB', (2200, 1600), 'white')
        draw = ImageDraw.Draw(canvas)
        subset = files[first:first + 4]
        for k, path in enumerate(subset):
            x = k % 2 * 1100
            y = k // 2 * 800
            draw.text((x + 10, y + 8), f'图{first + k + 1:02d}  {FIGURE_NAMES[path.stem]}', fill='black', font=font)
            picture = Image.open(path).convert('RGB')
            picture.thumbnail((1080, 740))
            canvas.paste(picture, (x + (1100 - picture.width) // 2, y + 45))
        name = f'contact_{first // 4 + 1}.png'
        canvas.save(dest / name)
        manifest.append(dict(sheet=name, figures=[p.name for p in subset]))
    write_json(dest / 'contact_manifest.json', manifest)
    print('contact sheets:', len(manifest))

# ==================== compile_complete_tables ====================
def compile_complete_tables_read(p):
    return json.loads(p.read_text(encoding='utf-8'))

def table(headers, rows):
    return ['| ' + ' | '.join(map(str, headers)) + ' |', '| ' + ' | '.join(['---'] * len(headers)) + ' |'] + ['| ' + ' | '.join(map(str, row)) + ' |' for row in rows]

def compile_tables():
    D = OUTPUT / '完整表格'
    D.mkdir(exist_ok=True)
    source = load_source()
    candidate = compile_complete_tables_read(OUTPUT / 'voltage_optimization/候选模型.json')
    lines = ['# B题完整表格汇编', '', '表1、表2采用当前shape_joint候选96格结果，与现有电压优化表保持一致；192格为新增网格核查，另列完整对照表。表3及题目中两个表4仍为旧冻结模型的既有控制计算，单列附录，不能作为新电压模型已完成第二至第四问的证据。', '', '所有冰体积分数均以控制体总体积为分母，无量纲；百分比列另行标明。各表保留原实验值和题目指定的8个时刻。相对误差由未舍入值计算；温度按题目采用摄氏度数值。显示0.000000只表示六位小数下舍入为零。', '']
    provenance = []
    for n, temp in enumerate((-20, -25), 1):
        p = OUTPUT / f'voltage_optimization/表{n}_电压优化候选.csv'
        with p.open(encoding='utf-8-sig', newline='') as f:
            rows = list(csv.reader(f))
        r = load_run(candidate['fine_results'][str(temp)]['run'])
        for given, rendered in zip(source['tables'][n - 1][1:], rows[1:]):
            t = float(given[0])
            v = float(np.interp(t, r['t'], r['V'][:, 0]))
            T = float(np.interp(t, r['t'], r['T'][:, 0]))
            ice = float(np.interp(t, r['t'], r['ice'][:, 0]))
            expected = [given[0], given[1], f'{v:.4f}', f'{100 * abs(v - float(given[1])) / abs(float(given[1])):.2f}', given[4], f'{T:.2f}', f'{100 * abs(T - float(given[4])) / abs(float(given[4])):.2f}', f'{ice:.6f}']
            assert list(map(str, expected)) == rendered, (expected, rendered)
        lines += [f'## 表{n} 一维单电池瞬态自冷启动模型预测与实验验证结果（{temp} ℃，当前96格候选）', ''] + table(rows[0], rows[1:]) + ['']
        provenance.append({'table': n, 'mesh': 96, 'run': candidate['fine_results'][str(temp)]['run'], 'source_csv': str(p.relative_to(OUTPUT)), 'sha256': hashlib.sha256(p.read_bytes()).hexdigest()})
    lines += ['注：这两表为联合校准参数下的回代结果；完整预测门槛尚未全部通过，不应表述为独立验证通过。零时刻严格从题设初温初始化，所以温度相对误差为零。', '']
    for n, temp in enumerate((-20, -25), 1):
        record = compile_complete_tables_read(OUTPUT / f'ice_analysis/grid192_{temp}.json')
        r = load_run(record['run'])
        original = source['tables'][n - 1]
        out = []
        for row in original[1:]:
            t = float(row[0])
            v = float(np.interp(t, r['t'], r['V'][:, 0]))
            T = float(np.interp(t, r['t'], r['T'][:, 0]))
            ice = float(np.interp(t, r['t'], r['ice'][:, 0]))
            out.append([row[0], row[1], f'{v:.4f}', f'{100 * abs(v - float(row[1])) / abs(float(row[1])):.2f}', row[4], f'{T:.2f}', f'{100 * abs(T - float(row[4])) / abs(float(row[4])):.2f}', f'{ice:.6f}'])
        lines += [f'## 核查表{n} 同参数192格完整复算（{temp} ℃）', ''] + table(original[0], out) + ['']
        provenance.append({'table': f'核查{n}', 'mesh': 192, 'run': record['run']})
    d = pd.read_csv(OUTPUT / 'ice_analysis/冰分数与水量逐时诊断.csv')
    for temp in (-20, -25):
        rows = []
        for t in (0, 5, 10, 15, 20, 25, 30, 35, 36.6):
            vals = []
            for mesh in (96, 192):
                part = d[(d.temp_C == temp) & (d.model == f'改进模型{mesh}格')]
                x = part.iloc[np.argmin(abs(part.time_s - t))]
                assert abs(x.time_s - t) < 1e-08
                vals.append(x)
            a, b = vals
            rows.append([f'{t:g}', f'{a.phi_max:.8f}', f'{b.phi_max:.8f}', f'{b.phi_max * 100:.5f}', f'{b.saturation_max * 100:.5f}', b.phi_layer if b.phi_max >= 1e-06 else '低于诊断阈值', b.saturation_layer if b.phi_max >= 1e-06 else '低于诊断阈值'])
        lines += [f'## 冰分数完整对照（{temp} ℃）', ''] + table(['时间/s', '96格最大冰体积分数', '192格最大冰体积分数', '192格最大冰体积分数/%', '192格最大孔隙冰饱和度/%', '192格冰体积分数峰值层', '192格冰饱和度峰值层'], rows) + ['']
    lines += ['36.6 s为实验记录终点补充行，不属于题目原表的8个指定时刻。诊断阈值为总体积冰分数10⁻⁶，非题设阈值或实验检出限。全域最大冰分数与最大冰饱和度可能位于不同网格。96→192格冰峰值分别变化36.48%和10.63%，尚不能称峰值已经收敛。', '', '## 附录：旧冻结模型表3及表4', '', '以下沿用既有控制计算，保留未找到可行解、离散模型成功及网格未收敛等限制。没有将本次96/192格单电池结果拼入这些旧控制结果。', '']
    old = (OUTPUT / 'B题原表填写/B题按题目要求填写的表格.md').read_text(encoding='utf-8')
    lines.append(old[old.index('## 表3 '):])
    lines += ['', '## 数据来源', '', '当前模型表1/2：outputs/voltage_optimization；冰量及192格核查：outputs/ice_analysis；旧控制模型表3/4：outputs/B题原表填写。所有运行编号及源文件哈希见同目录“表格来源与核查.json”。']
    (D / 'B题完整表格_当前候选与网格核查.md').write_text('\n'.join(lines) + '\n', encoding='utf-8')
    (D / '表格来源与核查.json').write_text(json.dumps({'original_source_sha256': source['sha256'], 'tables': provenance, '96_grid_table_cells_match_saved_csv': True, 'source_model_selection_unchanged': True, 'old_control_tables_copied_verbatim': True}, ensure_ascii=False, indent=2), encoding='utf-8')
    print('Complete tables written and 96-grid values independently rechecked.')
    print((D / 'B题完整表格_当前候选与网格核查.md').read_text(encoding='utf-8').split('## 核查表1')[1].split('## 冰分数完整对照')[0])

# ==================== test_physics ====================
def test_geometry_and_units():
    m = Model(Config())
    assert sum(m.g.dx) == pytest.approx(0.0003267)
    assert np.dot(m.g.C, m.g.tdx) == pytest.approx(6127.47966)
    s = Model(Config(cells=5, fixture='H2'))
    assert np.dot(s.g.C, s.g.tdx) == pytest.approx(109637.3983)
    assert sum(s.g.tdx) == pytest.approx(0.0416335)
    assert AREA * 10000.0 == 25

def test_exact_saturation_and_enthalpy_inversion():
    m = Model(Config())
    y = m.initial()
    z, H, _ = m.unpack(y)
    z[0, 0, m.g.clc] = 100
    z[1, 0, m.g.clc] = 20
    T = np.full(m.g.nt, 253.15)
    mv, ml, eg = partition(z[0], z[1], T[m.g.mea_indices], m.g.eps)
    h = m.g.C * (T - TM)
    h[m.g.mea_indices] += (2000 * mv + 4182 * ml + 2050 * z[1] + 4182 * z[2]) * -20 + LV * mv - LF * z[1]
    y[m.ih:m.il] = h / m.g.C
    s = m.phases(y)
    assert np.max(abs(s['T'] - T)) < 1e-07
    assert np.max(abs(s['mv'] + s['ml'] - z[0])) < 1e-12
    assert np.min(eg[:, m.g.porous]) > 0

def test_conservative_residual_at_nonuniform_state():
    m = Model(Config(cells=5, fixture='H2', outlet='drain'))
    y = m.initial()
    z, _, _ = m.unpack(y)
    z[0, :, m.g.clc] = 2
    z[1, :, m.g.clc] = 1
    d = m.rhs(0, y, 0.1, np.array([0.1, 0.2, 0.3, 0.4, 0.5]))
    dz = d[:m.nw].reshape(4, m.c, m.n)
    water = np.sum(dz[:3] * m.g.dx, axis=(0, 2))
    assert np.max(abs(water - d[m.il:m.il + m.c])) < 1e-13
    thermal = np.dot(d[m.ih:m.il] * m.g.C, m.g.tdx)
    assert thermal == pytest.approx(d[-2], rel=1e-12, abs=1e-07)
    detail = m.rhs(0, y, 0.1, np.array([0.1, 0.2, 0.3, 0.4, 0.5]), details=True)
    for c in range(m.c):
        for layer in range(5):
            ix = np.flatnonzero(m.g.layer == layer)
            left, right = (ix[0], ix[-1] + 1)
            external = sum((f[c, left] - f[c, right] for name, f in detail['fluxes'].items() if name != 'gas'))
            generated = MW * 1000 / (2 * F) if layer == 3 else 0.0
            local = np.sum(dz[:3, c, ix] * m.g.dx[ix], axis=None)
            assert local == pytest.approx(external + generated, abs=1e-13)

def test_liquid_flux_dry_equal_pressure_and_drain():
    args = dict(T=np.array([253.15] * 2), eps=np.array([0.8] * 2), K=np.array([6.2e-12] * 2), angle=np.array([110.0] * 2), dx=np.array([1e-05] * 2), ice=np.zeros(2))
    assert np.all(liquid_flux(np.zeros(2), **args, drain=True) == 0)
    assert liquid_flux(np.array([100.0, 100.0]), **args)[1] == 0
    assert liquid_flux(np.array([100.0, 0.0]), **args)[1] > 0
    assert liquid_flux(np.array([0.0, 100.0]), **args)[1] < 0
    assert liquid_flux(np.array([100.0, 100.0]), **args, drain=True, side=0)[0] < 0
    blocked = {**args, 'ice': np.array([0.0, 920 * 0.8])}
    assert abs(liquid_flux(np.array([100.0, 0.0]), **blocked)[1]) < 1e-15

def test_plate_flux_cancels_and_sealed_energy():
    m = Model(Config(cells=5, fixture='H2', h=0))
    y = m.initial(Tfield=np.linspace(250, 280, m.g.nt))
    d = m.rhs(0, y, 0, np.zeros(5))
    assert np.dot(d[m.ih:m.il] * m.g.C, m.g.tdx) == pytest.approx(d[-2], abs=1e-07)

def test_fixed_charge_and_current_limits():
    p = Protocol('fixed')
    t = np.linspace(0, 60, 6001)
    assert np.trapezoid([p.j(ti) for ti in t], t) == pytest.approx(9)
    assert 0.3 * (290 / 3) - 9 == pytest.approx(20)

def test_initial_temperature_is_not_fitted():
    m = Model(Config(T0=248.15))
    y = m.initial()
    assert np.max(abs(m.phases(y)['T'] - 248.15)) < 1e-10
    assert np.max(abs(m.phases(y)['lam'][:, m.g.pem] - 3)) < 1e-10

def test_colored_jacobian_matches_independent_directional_difference():
    m = Model(Config(cells=5, fixture='H2'))
    y = m.initial()
    z, _, _ = m.unpack(y)
    z[0, :, m.g.porous] = 2
    z[1, :, m.g.porous] = 1
    u = np.arange(1, 6) * 0.1
    jac = m.jacobian(0, y, 0.1, u)
    rng = np.random.default_rng(7)
    direction = np.zeros(m.size)
    direction[m.active_columns] = rng.normal(size=len(m.active_columns))
    h = 1e-06
    direct = (m.rhs(0, y + h * direction, 0.1, u) - m.rhs(0, y - h * direction, 0.1, u)) / (2 * h)
    estimated = jac @ direction
    ix = np.r_[m.ih + np.flatnonzero(m.g.names == 'BP'), m.ih + np.flatnonzero(m.g.names == 'endplate')]
    assert np.linalg.norm((estimated - direct)[ix]) / np.linalg.norm(direct[ix]) < 0.0001
    assert len(m._jac_groups) < len(m.active_columns) / 2

# ==================== test_voltage_closure ====================
def test_uniform_cl_dissipation_resistance_is_mesh_independent():
    expected = 0.2 * 10000.0 * (THICKNESS[1] + THICKNESS[3]) / (3 * conductivity(253.15, 3) * 0.3 ** 1.5)
    for mesh in [(1, 1, 1, 1, 1), (3, 2, 3, 3, 3), (8, 8, 8, 8, 8)]:
        cfg = Config(mesh=mesh, cl_proton_loss=True)
        m = Model(cfg)
        s = m.phases(m.initial())
        new = m.voltage(s, 0.2)
        old = Model(cfg.changed(cl_proton_loss=False)).voltage(s, 0.2)
        assert new['cl_ohmic'][0] == pytest.approx(expected, rel=1e-12)
        assert old['V'][0] - new['V'][0] == pytest.approx(expected)
        assert m.voltage(s, 0)['cl_ohmic'][0] == 0

def test_hydration_factor_limits_and_reference_normalization():
    for gamma in (0.0, 3.0):
        m = Model(Config(cathode_hydration_exponent=gamma))
        s = m.phases(m.initial())
        assert m.voltage(s, 0.1)['active_fraction'][0] == pytest.approx((3 / 14) ** gamma)
        s['lam'][:, m.g.clc] = 28.0
        assert m.voltage(s, 0.1)['active_fraction'][0] == pytest.approx(1.0)

@pytest.mark.parametrize('equilibrium', ['ice_reference', 'supercooled_liquid'])
def test_extended_closure_still_closes_energy_and_water(equilibrium):
    m = Model(Config(cl_proton_loss=True, cathode_hydration_exponent=3, j0=15, vapor_equilibrium=equilibrium))
    y = m.initial()
    d = m.rhs(0, y, 0.1, np.zeros(1))
    dz = d[:m.nw].reshape(4, m.c, m.n)
    assert np.sum(dz[:3] * m.g.dx) == pytest.approx(d[m.il], abs=1e-12)
    assert np.dot(d[m.ih:m.il] * m.g.C, m.g.tdx) == pytest.approx(d[-2], abs=1e-07)

def test_supercooled_flash_conserves_water_and_uses_liquid_saturation():
    w = np.array([10.0])
    ice = np.array([1.0])
    eps = np.array([0.4])
    T = np.array([253.15])
    mv, ml, eg = partition(w, ice, T, eps, True)
    assert mv + ml == pytest.approx(w)
    assert mv / eg == pytest.approx(0.018 * psat(T, liquid=True) / (8.314462618 * T))
    assert mv[0] > partition(w, ice, T, eps, False)[0][0]

# ==================== 四问统一结果整理与运行检查 ====================
def export_results(question):
    """按题号发布已有结果；底层全精度轨迹仍保存在work目录。"""
    import shutil
    root = _project_root()
    out = _project_storage()/'outputs'
    dest = root/'results'/question
    dest.mkdir(parents=True, exist_ok=True)
    copied = []
    def copy_file(source, relative):
        if not source.is_file():
            return
        target = dest/relative
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, target)
        copied.append({'file': str(target.relative_to(dest)),
                       'source': str(source.relative_to(root)),
                       'sha256': hashlib.sha256(source.read_bytes()).hexdigest()})
    subdirs = ('calibration','voltage_optimization','ice_analysis') if question=='q1' else ()
    for folder in subdirs:
        for source in (out/folder).rglob('*'):
            if source.is_file(): copy_file(source, Path(folder)/source.relative_to(out/folder))
    extra = {
        'q1': ('selected_model','mesh_','preflight_','identifiability','parameter_','sensitivity','water_','slope_','temperature_slope','failure_diagnostics','cpu_profile'),
        'q2': ('q2_',),
        'q3': ('q3_','zero_heat_','energy_lower_bound','load_qualification','horizon_extension'),
        'q4': ('q4_','precool_','vref','controller_latency','control_robustness','extended_control_robustness')
    }[question]
    for source in out.iterdir():
        if source.is_file() and source.name.startswith(extra):copy_file(source, Path('计算记录')/source.name)
    for source in (out/'figures').glob('*'):
        if source.stem.lower().startswith(question):
            copy_file(source,Path('旧模型对照图' if question=='q1' else '图')/source.name)
    if question=='q1':
        for folder in ('voltage_optimization','ice_analysis'):
            for source in (out/folder).iterdir():
                if source.suffix in ('.png','.pdf','.svg'):copy_file(source,Path('图')/source.name)
        primary=[out/'voltage_optimization/表1_电压优化候选.csv',out/'voltage_optimization/表2_电压优化候选.csv']
    else:
        filename={'q2':'表3_问题2.csv','q3':'表4_问题3.csv','q4':'表4_问题4.csv'}[question]
        primary=[out/'B题原表填写'/filename]
    for source in (out/'tables').glob('*'+question.upper()+'*.csv'):
        copy_file(source,Path('计算表格')/source.name)
    text=[f'# {question.upper()} 结果', '',
          '本目录为按题号整理的结果副本；全精度计算记录保存在work/outputs，完整状态轨迹保存在work/runs。', '',
          'Q1保留当前电压优化候选及192格核查，Q2—Q4沿用旧冻结模型。模型分支未混用，已有未通过项与未找到可行解的结论保留。', '',
          '表格目录保留本次整理时已有的题设格式表；计算表格目录同步求解器生成的结构化表。后续重新搜索后，应以计算记录和计算表格中的运行编号为准，题设格式表及其注释属于既有定稿快照。', '']
    for source in primary:
        copy_file(source,Path('表格')/source.name)
        if source.exists():
            with source.open(encoding='utf-8-sig',newline='') as f: rows=list(csv.reader(f))
            if rows:
                text += ['## '+source.stem,'','| '+' | '.join(rows[0])+' |','| '+' | '.join(['---']*len(rows[0]))+' |']
                text += ['| '+' | '.join(row)+' |' for row in rows[1:]]
                text += ['']
    original=out/'B题原表填写/B题按题目要求填写的表格.md'
    if original.exists() and question!='q1':
        blocks=original.read_text(encoding='utf-8').split('## ')
        terms={'q2':'表3 ','q3':'表4 不同辅助','q4':'表4 不同预冷'}
        notes=next((part for part in blocks if part.startswith(terms[question])),None)
        if notes:text += ['## 原表完整注释与适用范围','','## '+notes]
    if question=='q1':
        text += ['冰体积分数是无量纲总体积比例；0.000000仅表示显示精度下舍入为零。',
                 '电压拟合为联合校准回代，不能表述为独立验证通过；96/192格冰峰值尚未收敛。','',
                 '![电压拟合曲线](图/电压拟合曲线.png)','',
                 '![冰峰值网格敏感性](图/冰峰值网格敏感性.png)']
    (dest/'结果说明.md').write_text('\n'.join(text)+'\n',encoding='utf-8')
    (dest/'文件清单.json').write_text(json.dumps(copied,ensure_ascii=False,indent=2),encoding='utf-8')
    print(f'{question}: 已整理 {len(copied)} 个结果文件 → {dest}')
    return dest


def self_check():
    """运行嵌入四个文件的原23项测试，不需要旧tests目录。"""
    paths=[str(Path(__file__).resolve().parent/(q+'.py')) for q in ('q1','q2','q3','q4')]
    result=subprocess.run([sys.executable,'-m','pytest',*paths,'-q','-o','python_functions=test_*','-p','no:cacheprovider'])
    if result.returncode:raise SystemExit(result.returncode)


def smoke_worker(case):
    """短时真实积分，用于检查独立导入和Windows多进程。"""
    model=Config(cells=case,fixture='H2' if case==5 else 'H1',horizon=.4,sample=.2,j0=.3,charge_limit=100)
    result=simulate(model,Protocol('constant',(.02,)))
    assert result['metrics']['last_valid_time']>=.4-1e-8
    assert result['metrics']['mass_residual']<1e-3
    assert result['metrics']['energy_residual']<1e-3
    return {'cells':case,'stop_reason':result['metrics']['stop_reason'],'mass_residual':result['metrics']['mass_residual'],'energy_residual':result['metrics']['energy_residual']}


def smoke_check(workers=2):
    with ProcessPoolExecutor(max_workers=workers) as pool:
        results=list(pool.map(smoke_worker,[1,5]))
    print(json.dumps(results,ensure_ascii=False,indent=2))
    return results

# ==================== 四问命令入口 ====================
def main(argv=None):
    parser=argparse.ArgumentParser(description='Q1：公共守恒模型、实验校准、电压与冰量分析')
    parser.add_argument('command',nargs='?',default='results',choices=['results','check','smoke','preflight','calibrate','optimize','replay','ice-refine','ice-analysis','figures','convergence','uncertainty','identifiability','tables'])
    parser.add_argument('--budget',type=int,default=12)
    parser.add_argument('--workers',type=int,default=2)
    parser.add_argument('--candidate',default='shape_joint')
    args=parser.parse_args(argv)
    if args.command=='check':self_check();return
    if args.command=='smoke':smoke_check(args.workers);return
    if args.command=='preflight':preflight()
    elif args.command=='calibrate':calibrate_all(args.workers,args.budget)
    elif args.command=='optimize':
        jobs=[('shape_joint',(-20,-25)),('shape_train_m20',(-20,)),('shape_train_m25',(-25,))]
        with ProcessPoolExecutor(max_workers=args.workers) as pool:
            futures=[pool.submit(optimize_voltage_fit_job,name,train,'liquid_cl',args.budget,True) for name,train in jobs]
            for future in as_completed(futures):future.result()
    elif args.command=='replay':
        verify_voltage_verify(args.candidate)
        verify_voltage_report(args.candidate)
    elif args.command=='ice-refine':
        with ProcessPoolExecutor(max_workers=args.workers) as pool:rows=list(pool.map(refine_ice_run,(-20,-25)))
        write_json(OUTPUT/'ice_analysis/grid192.json',rows)
    elif args.command=='ice-analysis':
        analyze_ice_main();report_ice_report();fit_ice_curve_main()
    elif args.command=='figures':
        figures();plot_voltage_fit_main();fit_ice_curve_main()
    elif args.command=='convergence':convergence(selected_config())
    elif args.command=='uncertainty':uncertainty(selected_config(),args.workers)
    elif args.command=='identifiability':identifiability(selected_config(),args.workers)
    elif args.command=='tables':make_tables();compile_tables()
    export_results('q1')

if __name__=='__main__':
    main()
